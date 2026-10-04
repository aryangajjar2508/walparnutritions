import os
import re
import math
import io
import json
import shutil
import difflib
from datetime import datetime
from pathlib import Path
import pandas as pd
from typing import List, Dict, Any, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from app.master_config import master_config_mgr

class BatchMasterEngine:
    _instance = None

    def __init__(self, rate_file_path: str = "rate avg.xlsx"):
        self.rate_file_path = rate_file_path
        self.formulations = self._load_formulations()
        self.liquid_packaging_costs = self._load_liquid_packaging_costs()
        self.capsule_pricing = self._load_capsule_pricing()
        self.sugar_cost_mapping = {"20%": 168, "30%": 252, "40%": 336, "50%": 420, "60%": 504}
        self.sorbitol_cost_mapping = {"20%": 216, "30%": 324, "40%": 432, "50%": 540, "60%": 648}
        self.predefined_liquid_other_ingredients = self._load_liquid_ingredients()
        self.ingredient_list = []
        self.rate_dict = {}
        self._load_rate_data()

    @classmethod
    def get_instance(cls):
        if cls._instance is None:
            cls._instance = BatchMasterEngine()
        return cls._instance

    def _load_rate_data(self):
        if os.path.exists(self.rate_file_path):
            try:
                rate_df = pd.read_excel(self.rate_file_path)
                if rate_df.empty or len(rate_df.columns) < 2:
                    rate_df = pd.read_excel(self.rate_file_path, header=None)
                if not rate_df.empty and len(rate_df.columns) >= 2:
                    col0 = rate_df.columns[0]
                    col1 = rate_df.columns[1]
                    self.ingredient_list = []
                    self.rate_dict = {}
                    for _, row in rate_df.iterrows():
                        ing_str = str(row[col0]).strip()
                        if not ing_str or ing_str.lower() in ["item", "ingredient", "ingredient / material", "none", "nan"]:
                            continue
                        try:
                            rate_val = float(row[col1])
                            self.ingredient_list.append(ing_str)
                            self.rate_dict[ing_str.lower()] = rate_val
                        except (ValueError, TypeError):
                            continue
                    print(f"[BatchMasterEngine] Loaded {len(self.ingredient_list)} ingredients from {self.rate_file_path}.")
            except Exception as e:
                print(f"Warning: Could not load {self.rate_file_path}: {e}")

    def _load_ai_cache(self) -> Dict[str, str]:
        cache_file = Path(__file__).parent / "data" / "ai_rate_cache.json"
        if cache_file.exists():
            try:
                with open(cache_file, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                return {}
        return {}

    def _save_ai_cache_match(self, query: str, matched_name: str):
        cache_file = Path(__file__).parent / "data" / "ai_rate_cache.json"
        cache = self._load_ai_cache()
        cache[query.strip().lower()] = matched_name.strip()
        try:
            cache_file.parent.mkdir(parents=True, exist_ok=True)
            with open(cache_file, "w", encoding="utf-8") as f:
                json.dump(cache, f, indent=2, ensure_ascii=False)
        except Exception as e:
            print(f"[BatchMasterEngine] Cache write error: {e}")

    def find_similar_ingredient_ai(self, query: str) -> Optional[str]:
        """
        Use Google Gemini AI API to match an OCR-extracted ingredient name
        to the single closest / equivalent ingredient from the new rate list.
        Results are cached in ai_rate_cache.json for instant repeated lookup.
        """
        if not query or len(query.strip()) < 2:
            return None
            
        clean_q = query.strip()
        q_lower = clean_q.lower()

        # Check knowledge cache first
        cache = self._load_ai_cache()
        if q_lower in cache:
            cached_name = cache[q_lower]
            if cached_name and cached_name.lower() in self.rate_dict:
                return cached_name

        try:
            import google.generativeai as genai
            from app.config import GEMINI_API_KEY
            genai.configure(api_key=GEMINI_API_KEY)
            model = genai.GenerativeModel("gemini-2.5-flash")

            catalog = self.ingredient_list
            if not catalog:
                catalog = list(self.rate_dict.keys())

            prompt = f"""You are a chief pharmaceutical formulation chemist at Walpar.
Match this OCR ingredient name: '{clean_q}' to the single most equivalent or chemically identical ingredient from our company pricing catalog:
{json.dumps(catalog)}

CRITICAL RULES:
1. If there is a direct match, common pharmaceutical synonym, salt form, or botanical equivalent (e.g. Pyridoxine HCL -> VITAMIN B6, 5-Hydroxytryptophan / Griffonia -> 5-HTP, Ascorbic Acid -> VITAMIN C PLAIN (ASCORBIC ACID), Ubiquinone / CoQ10 -> CO ENZYME Q-10, Gotu Kola -> CENTELLA ASIATICA, Withania Somnifera -> ASHWAGANDHA DRY EXTRACT, Methylcobalamin -> METHYL COBALAMIN, Threonate -> MAGNESIUM L-THERONATE, Cholecalciferol -> CHOLECALCIFEROL (VIT D3) VIAL), return that EXACT catalog name string.
2. If genuinely unrelated to any listed catalog item, return 'NONE'.
3. Return ONLY a single JSON object: {{"matched": "EXACT_CATALOG_NAME_OR_NONE"}}"""

            resp = model.generate_content(prompt)
            txt = resp.text.strip().replace("```json", "").replace("```", "").strip()
            data = json.loads(txt)
            matched = data.get("matched", "").strip()

            if matched and matched.upper() != "NONE" and matched.lower() in self.rate_dict:
                self._save_ai_cache_match(q_lower, matched)
                print(f"[BatchMasterEngine] AI matched '{clean_q}' -> '{matched}'")
                return matched
        except Exception as e:
            print(f"[BatchMasterEngine] AI similarity lookup error for '{clean_q}': {e}")

        return None

    def search_rates(self, query: str, limit: int = 15) -> List[Dict[str, Any]]:
        q = query.strip().lower()
        results = []
        if not q:
            for item in self.ingredient_list[:limit]:
                results.append({"name": item, "rate": self.rate_dict.get(item.lower(), 0.0)})
            return results

        exact = []
        startswith = []
        contains = []
        for item in self.ingredient_list:
            item_lower = item.lower()
            if item_lower == q:
                exact.append(item)
            elif item_lower.startswith(q):
                startswith.append(item)
            elif q in item_lower:
                contains.append(item)

        matches = (exact + startswith + contains)[:limit]
        for item in matches:
            results.append({"name": item, "rate": self.rate_dict.get(item.lower(), 0.0)})
        return results

    def get_rate_info(self, name: str) -> Dict[str, Any]:
        """
        Check if an ingredient is available in rate avg.xlsx.
        Returns:
            {
                "available": bool,
                "rate": float,
                "matched_name": str,
                "is_exact": bool
            }
        """
        if not name:
            return {"available": False, "rate": 0.0, "matched_name": "", "is_exact": False}

        raw_name = str(name).strip()
        clean_name = raw_name.lower()
        norm_name = re.sub(r'[\'\"\,]', '', clean_name).strip()
        norm_hyphen = norm_name.replace('-', ' ')

        def _is_capsule_shell(k: str, v: float) -> bool:
            if v < 1.5:
                return True
            k_low = k.lower()
            if any(tag in k_low for tag in ['veg', 'non veg', 'ct/ct', 'hpmc', "''00''", "''0''", 'size "0"', "size '0'", 'size "00"']):
                return True
            return False

        active_rates = {k: v for k, v in self.rate_dict.items() if not _is_capsule_shell(k, v)}

        # 1. Exact match in active rate database
        if clean_name in active_rates:
            return {"available": True, "rate": float(active_rates[clean_name]), "matched_name": clean_name, "is_exact": True}
        if norm_name in active_rates:
            return {"available": True, "rate": float(active_rates[norm_name]), "matched_name": norm_name, "is_exact": True}
        if norm_hyphen in active_rates:
            return {"available": True, "rate": float(active_rates[norm_hyphen]), "matched_name": norm_hyphen, "is_exact": True}

        # 2. Spelling Normalization (carnitine <-> cernitine, blueberry <-> blue berry, coq10 <-> co enzyme q 10)
        norm_spelling = clean_name
        norm_spelling = re.sub(r'\bblueberry\b', 'blue berry', norm_spelling)
        norm_spelling = re.sub(r'\bcarnitine\b', 'cernitine', norm_spelling)
        norm_spelling = re.sub(r'\bco\s*q\s*10\b|\bcoq10\b', 'co enzyme q 10', norm_spelling)
        norm_spelling = re.sub(r'\bcoenzyme\s*q\s*10\b', 'co enzyme q 10', norm_spelling)
        norm_spelling = re.sub(r'\bext\b', 'extract', norm_spelling)

        if norm_spelling in active_rates:
            return {"available": True, "rate": float(active_rates[norm_spelling]), "matched_name": norm_spelling, "is_exact": True}

        # 3. Known pharmaceutical synonyms in rate avg.xlsx (new rate list)
        synonyms = {
            "5-htp": ["5-htp"],
            "griffonia": ["5-htp"],
            "5-hydroxytryptophan": ["5-htp"],
            "hydroxytryptophan": ["5-htp"],
            "l-glutathione": ["l-glutathione reduced", "opitac glutathione", "liposomol glutathione 40%"],
            "glutathione": ["l-glutathione reduced", "opitac glutathione"],
            "l glutathione": ["l-glutathione reduced", "opitac glutathione"],
            "glutathion": ["l-glutathione reduced", "opitac glutathione"],
            "vitamin c": ["vitamin c plain (ascorbic acid)", "vitamin c coated"],
            "ascorbic acid": ["vitamin c plain (ascorbic acid)", "vitamin c coated"],
            "vitamin d3": ["vitamin d3 1 lakh iu", "vitamin d3 veg source", "cholecalciferol (vit d3) vial", "cholecalciferol ip", "vitamin d3 5lakh iu"],
            "cholecalciferol": ["cholecalciferol (vit d3) vial", "cholecalciferol ip", "vitamin d3 1 lakh iu"],
            "vitamin e": ["vitamin e 50% powder", "vitamin e acetate liquid"],
            "tocopherol": ["vitamin e 50% powder", "vitamin e acetate liquid"],
            "vitamin b12": ["methyl cobalamin"],
            "methylcobalamin": ["methyl cobalamin"],
            "cyanocobalamin": ["methyl cobalamin"],
            "cobalamin": ["methyl cobalamin"],
            "vitamin b6": ["vitamin b6"],
            "pyridoxine": ["vitamin b6"],
            "pyridoxal": ["vitamin b6"],
            "vitamin b1": ["vitamin b1 mono"],
            "thiamine": ["vitamin b1 mono"],
            "thiamin": ["vitamin b1 mono"],
            "vitamin b2": ["vitamin b2 plain"],
            "riboflavin": ["vitamin b2 plain"],
            "vitamin b3": ["niacinamide", "nicotinamide adenine dinucleotide nad"],
            "niacinamide": ["niacinamide"],
            "nicotinamide": ["niacinamide"],
            "vitamin b5": ["vitamin b5"],
            "pantothenic": ["vitamin b5"],
            "pantothenate": ["vitamin b5"],
            "vitamin k": ["vitamin k-27"],
            "vitamin k2": ["vitamin k-27"],
            "k2-7": ["vitamin k-27"],
            "k2 7": ["vitamin k-27"],
            "zinc": ["zinc bis glycinate", "zinc gluconate", "zinc picolinate", "zinc sulphate mono", "zinc citrate", "zinc oxide"],
            "iron": ["liposomal iron", "liposomal iron 20%", "iron carbonyl"],
            "folic acid": ["folic acid", "l-methyl folate"],
            "folate": ["l-methyl folate", "folic acid"],
            "methylfolate": ["l-methyl folate"],
            "biotin": ["d-biotin"],
            "d-biotin": ["d-biotin"],
            "curcumin": ["curcumin", "curcumin 95%", "curcumin dry extract", "curcumin 10% extract"],
            "turmeric": ["curcumin", "curcumin 95%"],
            "piperine": ["piperine"],
            "black pepper": ["piperine"],
            "bioperine": ["piperine"],
            "ashwagandha": ["ashwagandha dry extract", "ksm-66 ashwagandha"],
            "withania": ["ashwagandha dry extract", "ksm-66 ashwagandha"],
            "somnifera": ["ashwagandha dry extract", "ksm-66 ashwagandha"],
            "melatonin": ["melatonin"],
            "l-theanine": ["l-theanine"],
            "theanine": ["l-theanine"],
            "glucosamine": ["glucosamine sulphate", "glucosamine hcl"],
            "chondroitin": ["chondroitin sulphate", "chondroitin sulphate 90%", "chondroitin sulphate 40%"],
            "msm": ["msm"],
            "methylsulfonylmethane": ["msm"],
            "magnesium acetyl taurate": ["magnesium acetyl taurate", "magnesium taurate"],
            "magnesium taurate": ["magnesium taurate", "magnesium acetyl taurate"],
            "magnesium bisglycinate": ["magnesium bis glycinate", "magnesium bis glycinate 22%", "magnesium bis glycinate 30%"],
            "magnesium bis-glycinate": ["magnesium bis glycinate"],
            "magnesium bis glycinate": ["magnesium bis glycinate"],
            "magnesium l threonate": ["magnesium l-theronate"],
            "magnesium l-threonate": ["magnesium l-theronate"],
            "magnesium threonate": ["magnesium l-theronate"],
            "magnesium l theronate": ["magnesium l-theronate"],
            "calcium carbonate": ["calcium carbonate", "calcium carbonate algaecal"],
            "calcium citrate": ["calcium citrate", "calcium citrate malate"],
            "coral grains": ["calcium carbonate"],
            "co enzyme q 10": ["co enzyme q-10"],
            "coq10": ["co enzyme q-10"],
            "co-q10": ["co enzyme q-10"],
            "ubiquinone": ["co enzyme q-10"],
            "astaxanthin": ["astaxanthin 10% powder"],
            "blueberry": ["blue berry dry extract"],
            "blue berry": ["blue berry dry extract"],
            "bilberry": ["bilberry dry ext"]
        }

        for syn_key, targets in synonyms.items():
            if syn_key in clean_name or syn_key in norm_hyphen or syn_key in norm_spelling:
                for target in targets:
                    if target in active_rates:
                        return {"available": True, "rate": float(active_rates[target]), "matched_name": target, "is_exact": False}

        # 4. Smart RapidFuzz Similarity Matching with collision guards
        try:
            from rapidfuzz import fuzz, process

            # Collision guards for compounds with easily confused names
            is_glutamine = "glutamin" in clean_name
            is_glutathione = "glutathion" in clean_name
            is_thiamine = "thiamin" in clean_name
            is_theanine = "theanin" in clean_name
            is_carnosine = "carnosine" in clean_name

            # Guard against cross-matching distinct berry types (bilberry != raspberry != cranberry != blueberry)
            berries = ['bilberry', 'blue berry', 'blueberry', 'cranberry', 'raspberry', 'rasberry', 'strawberry', 'chaste berry', 'chasteberry', 'elderberry']
            q_berry = next((b for b in berries if b in norm_spelling), None)

            candidate_keys = list(active_rates.keys())
            if q_berry:
                if q_berry in ['blue berry', 'blueberry']:
                    candidate_keys = [k for k in candidate_keys if 'blue berry' in k or 'blueberry' in k]
                elif q_berry in ['raspberry', 'rasberry']:
                    candidate_keys = [k for k in candidate_keys if 'raspberry' in k or 'rasberry' in k]
                elif q_berry in ['bilberry']:
                    candidate_keys = [k for k in candidate_keys if 'bilberry' in k]
                elif q_berry in ['cranberry']:
                    candidate_keys = [k for k in candidate_keys if 'cranberry' in k]

            # Try token_set_ratio first (handles sub-specifications like 'extract', 'powder', '%')
            best_set = process.extractOne(norm_spelling, candidate_keys, scorer=fuzz.token_set_ratio)
            if best_set and best_set[1] >= 75.0:
                target_key = best_set[0]
                # Safety checks
                if is_glutamine and "glutathion" in target_key:
                    pass
                elif is_glutathione and "glutamin" in target_key:
                    pass
                elif is_thiamine and "theanin" in target_key:
                    pass
                elif is_theanine and "thiamin" in target_key:
                    pass
                elif is_carnosine and "cernitin" in target_key:
                    pass
                else:
                    return {"available": True, "rate": float(active_rates[target_key]), "matched_name": target_key, "is_exact": False}

            # Try token_sort_ratio as secondary
            best_sort = process.extractOne(norm_spelling, candidate_keys, scorer=fuzz.token_sort_ratio)
            if best_sort and best_sort[1] >= 75.0:
                target_key = best_sort[0]
                if not ((is_glutamine and "glutathion" in target_key) or (is_glutathione and "glutamin" in target_key)):
                    return {"available": True, "rate": float(active_rates[target_key]), "matched_name": target_key, "is_exact": False}

        except Exception as fuzz_err:
            pass

        # 5. Token Subset fallback
        for k, v in active_rates.items():
            k_clean = k.replace("'", "").replace('"', '').strip()
            tokens = [t for t in re.split(r'[\s\-\(\)\/]+', norm_spelling) if len(t) > 3]
            k_tokens = [t for t in re.split(r'[\s\-\(\)\/]+', k_clean) if len(t) > 3]
            if tokens and all(t in k_clean for t in tokens):
                return {"available": True, "rate": float(v), "matched_name": k, "is_exact": False}
            if k_tokens and all(t in norm_spelling for t in k_tokens):
                return {"available": True, "rate": float(v), "matched_name": k, "is_exact": False}

        # 6. AI Fallback: Google Gemini API for similar/equivalent name resolution
        try:
            ai_matched = self.find_similar_ingredient_ai(clean_name)
            if ai_matched and ai_matched.lower() in active_rates:
                return {
                    "available": True,
                    "rate": float(active_rates[ai_matched.lower()]),
                    "matched_name": ai_matched,
                    "is_exact": False,
                    "is_ai_matched": True
                }
        except Exception as ai_err:
            print(f"[BatchMasterEngine] AI fallback error for '{clean_name}': {ai_err}")

        # Not found in database
        return {"available": False, "rate": 0.0, "matched_name": "", "is_exact": False}

    def match_rate_for_ingredient(self, name: str) -> float:
        return float(self.get_rate_info(name).get("rate", 0.0))

    def batch_check_rates(self, names: List[str]) -> Dict[str, Dict[str, Any]]:
        results = {}
        for n in names:
            results[n] = self.get_rate_info(n)
        return results

    def _load_formulations(self):
        return {
            100: {"batch_kg": 0.01, "active_mg_per_tablet": 15, "ingredients": [
                ("Starch DCP granules", "DCP", "Powder", 0.003, 70, "B. Pest"),
                ("Starch paste", "Starch", "Powder", 0.0007, 45, "B. Pest"),
                ("Talcum", "Talcum", "Powder", 0.0007, 25, "C. Lubrication"),
                ("Magnesium Stearate", "Magnesium Stearate", "Powder", 0.0005, 140, "C. Lubrication"),
                ("SSG", "SSG", "Powder", 0.0004, 72, "C. Lubrication"),
                ("MCCP", "MCCP", "Powder", 0.0004, 450, "C. Lubrication"),
                ("Aerosil", "Aerosil", "Powder", 0.0001, 425, "C. Lubrication"),
                ("COATING - IRON OXIDE RED", "Coating Powder", "Powder", 0.0003, 950, "D. Coating"),
                ("IPA", "IPA", "Liquid", 0.0025, 85, "D. Coating"),
                ("MDC", "MDC", "Liquid", 0.0035, 58, "D. Coating")
            ]},
            400: {"batch_kg": 0.04, "active_mg_per_tablet": 150, "ingredients": [
                ("Starch DCP granules", "DCP", "Powder", 0.012, 70, "B. Pest"),
                ("Starch paste", "Starch", "Powder", 0.002, 45, "B. Pest"),
                ("Talcum", "Talcum", "Powder", 0.004, 25, "C. Lubrication"),
                ("Magnesium Stearate", "Magnesium Stearate", "Powder", 0.003, 140, "C. Lubrication"),
                ("SSG", "SSG", "Powder", 0.002, 72, "C. Lubrication"),
                ("MCCP", "MCCP", "Powder", 0.002, 450, "C. Lubrication"),
                ("Aerosil", "Aerosil", "Powder", 0.0005, 425, "C. Lubrication"),
                ("COATING - IRON OXIDE RED", "Coating Powder", "Powder", 0.0015, 950, "D. Coating"),
                ("IPA", "IPA", "Liquid", 0.012, 85, "D. Coating"),
                ("MDC", "MDC", "Liquid", 0.018, 58, "D. Coating")
            ]},
            900: {"batch_kg": 0.09, "active_mg_per_tablet": 337.5, "ingredients": [
                ("Starch DCP granules", "DCP", "Powder", 0.027, 70, "B. Pest"),
                ("Starch paste", "Starch", "Powder", 0.0045, 45, "B. Pest"),
                ("Talcum", "Talcum", "Powder", 0.009, 25, "C. Lubrication"),
                ("Magnesium Stearate", "Magnesium Stearate", "Powder", 0.00675, 140, "C. Lubrication"),
                ("SSG", "SSG", "Powder", 0.0045, 72, "C. Lubrication"),
                ("MCCP", "MCCP", "Powder", 0.0045, 450, "C. Lubrication"),
                ("Aerosil", "Aerosil", "Powder", 0.001125, 425, "C. Lubrication"),
                ("COATING - IRON OXIDE RED", "Coating Powder", "Powder", 0.003, 950, "D. Coating"),
                ("IPA", "IPA", "Liquid", 0.025, 85, "D. Coating"),
                ("MDC", "MDC", "Liquid", 0.035, 58, "D. Coating")
            ]},
            1000: {"batch_kg": 0.10, "active_mg_per_tablet": 420, "ingredients": [
                ("Starch DCP granules", "DCP", "Powder", 0.030, 70, "B. Pest"),
                ("Starch paste", "Starch", "Powder", 0.007, 45, "B. Pest"),
                ("Talcum", "Talcum", "Powder", 0.007, 25, "C. Lubrication"),
                ("Magnesium Stearate", "Magnesium Stearate", "Powder", 0.005, 140, "C. Lubrication"),
                ("SSG", "SSG", "Powder", 0.004, 72, "C. Lubrication"),
                ("MCCP", "MCCP", "Powder", 0.004, 450, "C. Lubrication"),
                ("Aerosil", "Aerosil", "Powder", 0.001, 425, "C. Lubrication"),
                ("COATING - IRON OXIDE RED", "Coating Powder", "Powder", 0.003, 950, "D. Coating"),
                ("IPA", "IPA", "Liquid", 0.025, 85, "D. Coating"),
                ("MDC", "MDC", "Liquid", 0.035, 58, "D. Coating")
            ]},
            1800: {"batch_kg": 0.18, "active_mg_per_tablet": 800, "ingredients": [
                ("Starch DCP granules", "DCP", "Powder", 0.054, 70, "B. Pest"),
                ("Starch paste", "Starch", "Powder", 0.010, 45, "B. Pest"),
                ("Talcum", "Talcum", "Powder", 0.012, 25, "C. Lubrication"),
                ("Magnesium Stearate", "Magnesium Stearate", "Powder", 0.009, 140, "C. Lubrication"),
                ("SSG", "SSG", "Powder", 0.007, 72, "C. Lubrication"),
                ("MCCP", "MCCP", "Powder", 0.007, 450, "C. Lubrication"),
                ("Aerosil", "Aerosil", "Powder", 0.002, 425, "C. Lubrication"),
                ("COATING - IRON OXIDE RED", "Coating Powder", "Powder", 0.006, 950, "D. Coating"),
                ("IPA", "IPA", "Liquid", 0.050, 85, "D. Coating"),
                ("MDC", "MDC", "Liquid", 0.070, 58, "D. Coating")
            ]}
        }

    def _load_capsule_pricing(self):
        return {
            "2": {"veg": 0.48, "non_veg": 0.18, "hpmc": 0.48, "gelatin": 0.18, "max_fill_mg": 150},
            "1": {"veg": 0.48, "non_veg": 0.18, "hpmc": 0.48, "gelatin": 0.18, "max_fill_mg": 250},
            "0": {"veg": 0.48, "non_veg": 0.18, "hpmc": 0.48, "gelatin": 0.18, "max_fill_mg": 450},
            "00": {"veg": 0.48, "non_veg": 0.18, "hpmc": 0.48, "gelatin": 0.18, "max_fill_mg": 650},
            "000": {"veg": 0.60, "non_veg": 0.30, "hpmc": 0.60, "gelatin": 0.30, "max_fill_mg": 1000},
            "DR": {"veg": 0.75, "hpmc": 0.75, "size": "0", "default": 0.75}
        }

    def _load_liquid_packaging_costs(self):
        return {
            "Bottle type": {
                "100ml": {"round brute": 2.25, "brute": 2.20, "dome": 2.30},
                "200ml": {"brute": 3.00, "micro brute": 3.60, "glo back": 3.50}
            },
            "ROPP type": {"quality": 0.55, "printed": 0.60},
            "measuring cup": {"yes": 0.25, "no": 0.0},
            "carton": {
                "100ml": {"uv dripp off": 2.30, "metallic": 3.50},
                "200ml": {
                    "uv dripp off": {"default": 2.45, "micro brute": 3.50},
                    "metallic": {"default": 4.00, "micro brute": 4.50}
                }
            },
            "label": {"chromo": 0.50, "metallic": 1.30},
            "shipper": {
                "100ml": {"100 nos": {"5ply": 50.0, "7ply": 60.0}},
                "200ml": {
                    "72 nos": {"5ply": 50.0, "7ply": 65.0},
                    "60 nos": {"5ply": 45.0, "7ply": 60.0}
                }
            },
            "accessories": {"yes": 0.50, "no": 0.0}
        }

    def _load_liquid_ingredients(self):
        cfg = master_config_mgr.get_config().get("liquid", {}).get("manufacturing", {})
        return [
            {"name": "Glycerin", "other": "Glycerin", "type": "Liquid", "qty_per_1000_bottles_kg": 4.0, "rate": cfg.get("glycerin_rate_per_kg", 180.0), "group": "B. Excipients"},
            {"name": "Tween 80", "other": "Tween 80", "type": "Liquid", "qty_per_1000_bottles_kg": 0.2, "rate": cfg.get("tween_80_rate_per_kg", 300.0), "group": "B. Excipients"},
            {"name": "Di Sodium EDTA", "other": "Di Sodium EDTA", "type": "Powder", "qty_per_1000_bottles_kg": 0.08, "rate": 250.0, "group": "B. Excipients"},
            {"name": "Menthol", "other": "Manthol", "type": "crystal", "qty_per_1000_bottles_kg": 0.013, "rate": 1400.0, "group": "B. Excipients"},
            {"name": "Sodium Citrate", "other": "Sodium Citrate", "type": "Powder", "qty_per_1000_bottles_kg": 0.2, "rate": 350.0, "group": "B. Excipients"},
            {"name": "Peppermint Flavour", "other": "Pipermint Flavour", "type": "Liquid", "qty_per_1000_bottles_kg": 0.1, "rate": 550.0, "group": "B. Excipients"},
            {"name": "Sweet Orange Flavour", "other": "Sweet orange flavor", "type": "Liquid", "qty_per_1000_bottles_kg": 0.096, "rate": 550.0, "group": "B. Excipients"},
            {"name": "Xanthan Gum - Transparent", "other": "Xantham Gum-transperent", "type": "Powder", "qty_per_1000_bottles_kg": 0.62, "rate": cfg.get("xanthan_gum_rate_per_kg", 1400.0), "group": "B. Excipients"},
            {"name": "AROSIL", "other": "AROSIL", "type": "Powder", "qty_per_1000_bottles_kg": 0.5, "rate": 400.0, "group": "B. Excipients"},
            {"name": "MCCP", "other": "MCCP", "type": "Powder", "qty_per_1000_bottles_kg": 1.0, "rate": 150.0, "group": "B. Excipients"},
            {"name": "Sodium Benzoate", "other": "Sodium Benzoate", "type": "Powder", "qty_per_1000_bottles_kg": 0.2, "rate": cfg.get("sodium_benzoate_rate_per_kg", 250.0), "group": "B. Excipients"},
            {"name": "Sodium Methyl Parabene", "other": "Sodium Methyl Parabene", "type": "Powder", "qty_per_1000_bottles_kg": 0.06, "rate": cfg.get("sodium_methyl_parabene_rate_per_kg", 500.0), "group": "B. Excipients"},
            {"name": "Sodium Propyl Parabene", "other": "Sodium Propyl Parabene", "type": "Powder", "qty_per_1000_bottles_kg": 0.05, "rate": cfg.get("sodium_propyl_parabene_rate_per_kg", 550.0), "group": "B. Excipients"}
        ]

    def get_default_excipients(self, product_type: str, tablet_size: int = 400) -> List[Dict[str, Any]]:
        excipients = []
        if product_type == "Tablet":
            form = self.formulations.get(tablet_size, self.formulations[400])
            for item in form["ingredients"]:
                excipients.append({
                    "name": item[0],
                    "other_name": item[1],
                    "type": item[2],
                    "qty_per_100": item[3],
                    "rate": item[4],
                    "group": item[5],
                    "selected": True
                })
        elif product_type == "Capsule":
            capsule_excipients = [
                ("Di-Calcium Phosphate", "Di-Calcium Phosphate", "Powder", 0.003, 70, "B. Excipients"),
                ("Talcum", "Talcum", "Powder", 0.0007, 25, "B. Excipients"),
                ("Sodium Methyl Parabene", "Sodium Methyl Parabene", "Powder", 0.0001, 100, "B. Excipients"),
                ("Sodium Propyl Parabene", "Sodium Propyl Parabene", "Powder", 0.000025, 25, "B. Excipients")
            ]
            for item in capsule_excipients:
                excipients.append({
                    "name": item[0],
                    "other_name": item[1],
                    "type": item[2],
                    "qty_per_100": item[3],
                    "rate": item[4],
                    "group": item[5],
                    "selected": True
                })
        elif product_type == "Liquid":
            for item in self.predefined_liquid_other_ingredients:
                excipients.append({
                    "name": item["name"],
                    "other_name": item["other"],
                    "type": item["type"],
                    "qty_per_1000": item["qty_per_1000_bottles_kg"],
                    "rate": item["rate"],
                    "group": item.get("group", "B. Excipients"),
                    "selected": True
                })
        return excipients

    def calculate_bmr(self, payload: Dict[str, Any]) -> Dict[str, Any]:
        """
        Takes configuration from user UI:
        - product_type: Tablet | Capsule | Liquid
        - quantity: total batch units
        - size_or_capsule: tablet mg (int), capsule size ('00','0','1','2'), or bottle size ('100 ml')
        - capsule_type: 'veg', 'non_veg', 'DR'
        - serving_size: '5 ml', '10 ml' (for liquid)
        - sugar_type: 'Sugar', 'Sorbitol', 'None'
        - sugar_percent: '20%', '30%', etc.
        - active_ingredients: list of {name, per_unit_mg, rate}
        - selected_excipients: list of {name, other_name, type, qty_per_100/qty_per_1000, rate, group, selected}
        - packaging: dict of packaging selections
        """
        product_type = payload.get("product_type", "Tablet")
        quantity = int(payload.get("quantity", 100000))
        # Enforce minimum batch size: 1,00,000 for Tablet and Capsule
        if product_type in ("Tablet", "Capsule") and quantity < 100000:
            quantity = 100000

        size_or_capsule = payload.get("size_or_capsule", 400)
        capsule_type = payload.get("capsule_type", "veg")
        serving_size = payload.get("serving_size", "5 ml")
        sugar_type = payload.get("sugar_type", "Sugar")
        sugar_percent = payload.get("sugar_percent", "20%")
        active_ingredients = payload.get("active_ingredients", [])
        selected_excipients = payload.get("selected_excipients", [])
        # Auto-select all default excipients if not provided
        if not selected_excipients:
            selected_excipients = self.get_default_excipients(product_type, size_or_capsule)
        # Calculate total active mg (especially for tablet auto-sizing & limit check)
        total_active_mg = 0.0
        for act in active_ingredients:
            raw_dose = float(act.get("dosage") or act.get("per_unit_mg") or act.get("mg_per_unit") or 0.0)
            u = str(act.get("unit") or "mg").lower().strip()
            name_lower = str(act.get("name") or "").lower()
            if u == "mcg":
                mg = raw_dose / 1000.0
            elif u == "g":
                mg = raw_dose * 1000.0
            elif u == "iu":
                if any(k in name_lower for k in ["vitamin a", "vit a", "retinol", "retinyl", "carotene"]):
                    mg = raw_dose * 0.0003
                elif any(k in name_lower for k in ["vitamin e", "vit e", "tocopher", "tocopheryl"]):
                    mg = raw_dose * 0.67
                else:
                    # Vitamin D family (D3 Cholecalciferol, D2 Ergocalciferol): 1 IU = 0.025 mcg = 0.000025 mg
                    mg = raw_dose * 0.000025
            else:
                mg = float(act.get("per_unit_mg", 0.0) or raw_dose)
            total_active_mg += mg

        auto_selected_size = None
        if product_type == "Tablet":
            # Strict rule: if it goes over 1800 mg, show please contact with us & don't calculate
            if total_active_mg > 1800.0:
                return {
                    "success": False,
                    "error": f"Total active weight is {total_active_mg:.1f} mg which exceeds the single-tablet limit of 1800 mg. Please contact with us. Formulation cannot be compressed into a single tablet.",
                    "contact_required": True,
                    "total_active_mg": round(total_active_mg, 2)
                }

            # Auto size selection:
            # if less than 50 mg then auto select 100 mg tablet
            # between 50 to 300 auto select 400 mg tablet
            # between 300 to 600 auto select 900 mg tablet
            # between 600 to 800 auto select 1000 mg tablet
            # between 800 to 1800 then select 1800 mg tablet
            if total_active_mg < 50.0:
                auto_selected_size = 100
            elif 50.0 <= total_active_mg <= 300.0:
                auto_selected_size = 400
            elif 300.0 < total_active_mg <= 600.0:
                auto_selected_size = 900
            elif 600.0 < total_active_mg <= 800.0:
                auto_selected_size = 1000
            else:
                auto_selected_size = 1800

            # Apply auto size unless user explicitly chose a manual override
            if not payload.get("manual_tablet_size_override") or not size_or_capsule:
                size_or_capsule = auto_selected_size
            else:
                try:
                    size_or_capsule = int(size_or_capsule)
                except Exception:
                    size_or_capsule = auto_selected_size

        elif product_type == "Capsule":
            # Strict rule: capsule max capacity is 1000 mg in Size 000. If > 1000 mg, show please contact with us & don't calculate
            if total_active_mg > 1000.0:
                return {
                    "success": False,
                    "error": f"Total active weight is {total_active_mg:.1f} mg which exceeds the maximum single-capsule capacity of 1000 mg (Size 000). Please contact with us. Formulation cannot be encapsulated into a single capsule.",
                    "contact_required": True,
                    "total_active_mg": round(total_active_mg, 2)
                }

            # Auto size selection based on maximum fill capacity:
            # <= 150 mg: Size "2"
            # 150 to 250 mg: Size "1"
            # 250 to 450 mg: Size "0"
            # 450 to 650 mg: Size "00"
            # 650 to 1000 mg: Size "000"
            if total_active_mg <= 150.0:
                auto_selected_size = "2"
            elif total_active_mg <= 250.0:
                auto_selected_size = "1"
            elif total_active_mg <= 450.0:
                auto_selected_size = "0"
            elif total_active_mg <= 650.0:
                auto_selected_size = "00"
            else:
                auto_selected_size = "000"

            if not payload.get("manual_capsule_size_override") or not size_or_capsule:
                size_or_capsule = auto_selected_size

        packaging = payload.get("packaging", {})
        items = []
        sr_no = 1
        total_active_cost = 0.0
        total_other_cost = 0.0
        total_batch_kg = 0.0

        # Group A: Active Ingredients
        if product_type == "Liquid":
            try:
                bottle_ml = float(str(size_or_capsule).replace(" ml", "").replace("ml", "").strip())
            except Exception:
                bottle_ml = 100.0
            try:
                srv_ml = float(str(serving_size).replace(" ml", "").replace("ml", "").strip())
            except Exception:
                srv_ml = 5.0
            servings_per_bottle = bottle_ml / srv_ml if srv_ml > 0 else 1.0
            total_servings = quantity * servings_per_bottle
        else:
            bottle_ml = 0
            srv_ml = 0
            servings_per_bottle = 1
            total_servings = quantity

        missing_rate_items = []
        for act in active_ingredients:
            name = act.get("name", "Active Ingredient")
            raw_dose = float(act.get("dosage") or act.get("per_unit_mg") or act.get("mg_per_unit") or 0.0)
            u = str(act.get("unit") or "mg").strip()
            u_lower = u.lower()
            name_lower = name.lower()

            if u_lower == "mcg":
                per_unit_mg = raw_dose / 1000.0
            elif u_lower == "g":
                per_unit_mg = raw_dose * 1000.0
            elif u_lower == "iu":
                if any(k in name_lower for k in ["vitamin a", "vit a", "retinol", "retinyl", "carotene"]):
                    per_unit_mg = raw_dose * 0.0003
                elif any(k in name_lower for k in ["vitamin e", "vit e", "tocopher", "tocopheryl"]):
                    per_unit_mg = raw_dose * 0.67
                else:
                    per_unit_mg = raw_dose * 0.000025
            else:
                per_unit_mg = float(act.get("per_unit_mg", 0.0) or raw_dose)
            
            rate_info = self.get_rate_info(name)
            user_rate = act.get("user_estimated_rate") or act.get("rate")
            is_user_estimated = bool(act.get("is_user_estimated"))

            if rate_info.get("available") and not is_user_estimated:
                rate = float(rate_info["rate"])
                is_rate_available = True
                matched_rate_name = rate_info["matched_name"]
            elif user_rate is not None and float(user_rate) > 0:
                rate = float(user_rate)
                is_rate_available = True
                is_user_estimated = True
                matched_rate_name = f"{name} (User Estimated)"
            elif rate_info.get("available"):
                rate = float(rate_info["rate"])
                is_rate_available = True
                matched_rate_name = rate_info["matched_name"]
            else:
                rate = 0.0
                is_rate_available = False
                matched_rate_name = ""
                missing_rate_items.append(name)

            if product_type == "Liquid":
                total_qty_mg = per_unit_mg * total_servings
                qty_kg = total_qty_mg / 1_000_000.0
                qty_g = qty_kg * 1000.0
                cost = qty_kg * rate
                if u_lower == "iu":
                    mg_fmt = f"{per_unit_mg:.6f}".rstrip('0').rstrip('.') if per_unit_mg < 0.001 else f"{per_unit_mg:.4f}"
                    unit_display = f"{raw_dose:.0f} IU / {srv_ml}ml ({mg_fmt} mg)"
                elif u_lower == "mcg":
                    unit_display = f"{raw_dose:.1f} mcg / {srv_ml}ml"
                else:
                    unit_display = f"{per_unit_mg:.2f} mg / {srv_ml}ml"
                cost_display = f"₹{cost/quantity:.2f}/btl" if quantity > 0 else ""
            else:
                total_qty_mg = per_unit_mg * quantity
                qty_kg = total_qty_mg / 1_000_000.0
                qty_g = qty_kg * 1000.0
                cost = qty_kg * rate
                if u_lower == "iu":
                    mg_fmt = f"{per_unit_mg:.6f}".rstrip('0').rstrip('.') if per_unit_mg < 0.001 else f"{per_unit_mg:.4f}"
                    unit_display = f"{raw_dose:.0f} IU ({mg_fmt} mg)"
                elif u_lower == "mcg":
                    unit_display = f"{raw_dose:.1f} mcg ({per_unit_mg:.3f} mg)"
                else:
                    unit_display = f"{per_unit_mg:.2f} mg"
                cost_display = f"₹{cost/quantity:.2f}/unit" if quantity > 0 else ""

            items.append({
                "sr_no": sr_no,
                "group": "A. Active Ingredients",
                "name": name,
                "other_name": act.get("other_name", name),
                "type": act.get("type", "Extract"),
                "qty_kg": round(qty_kg, 6),
                "qty_g": round(qty_g, 3),
                "qty_mg": round(total_qty_mg, 2),
                "qty_per_unit": unit_display,
                "rate": round(rate, 2),
                "cost": round(cost, 2),
                "is_rate_available": is_rate_available,
                "is_user_estimated": is_user_estimated,
                "rate_matched_name": matched_rate_name,
                "unit_cost_display": cost_display
            })
            total_active_cost += cost
            total_batch_kg += qty_kg
            sr_no += 1

        # Group B, C, D: Excipients
        if product_type == "Liquid":
            scale_factor = quantity / 1000.0
            st_lower = str(sugar_type).strip().lower()
            is_sugar_free = any(k in st_lower for k in ["sorbitol", "sugar free", "without sugar", "sugar-free"])

            for exc in selected_excipients:
                if not exc.get("selected", True):
                    continue
                name = exc.get("name", "")
                name_low = name.lower()

                # Mutually exclusive rule: either sugar or sorbitol, never both
                if is_sugar_free:
                    # Sugar-free selected -> exclude any Sugar ingredients
                    if "sugar" in name_low and "sorbitol" not in name_low:
                        continue
                else:
                    # Sugar base selected -> exclude Sorbitol ingredients
                    if "sorbitol" in name_low:
                        continue

                qty_per_1000 = float(exc.get("qty_per_1000", 0.0))
                rate = float(exc.get("rate", 0.0))
                qty_kg = qty_per_1000 * scale_factor
                qty_g = qty_kg * 1000.0
                qty_mg = qty_g * 1000.0
                cost = qty_kg * rate
                qty_per_serving_mg = (qty_mg / total_servings) if total_servings > 0 else 0.0

                items.append({
                    "sr_no": sr_no,
                    "group": exc.get("group", "B. Excipients"),
                    "name": name,
                    "other_name": exc.get("other_name", name),
                    "type": exc.get("type", "Powder"),
                    "qty_kg": round(qty_kg, 6),
                    "qty_g": round(qty_g, 3),
                    "qty_mg": round(qty_mg, 2),
                    "qty_per_unit": f"{qty_per_serving_mg:.2f} mg / {srv_ml}ml",
                    "rate": round(rate, 2),
                    "cost": round(cost, 2),
                    "unit_cost_display": ""
                })
                total_other_cost += cost
                total_batch_kg += qty_kg
                sr_no += 1

            # DM Water
            total_volume_L = (quantity * bottle_ml) / 1000.0
            dm_water_cost = total_volume_L * 1.00  # ₹1 / L
            items.append({
                "sr_no": sr_no,
                "group": "B. Base Liquid",
                "name": "D.M. Water",
                "other_name": "Purified Water",
                "type": "Liquid",
                "qty_kg": round(total_volume_L, 3),
                "qty_g": round(total_volume_L * 1000.0, 1),
                "qty_mg": round(total_volume_L * 1_000_000.0, 0),
                "qty_per_unit": f"{bottle_ml:.1f} ml / btl",
                "rate": 1.00,
                "cost": round(dm_water_cost, 2),
                "unit_cost_display": "₹1.00/L"
            })
            total_other_cost += dm_water_cost
            total_batch_kg += total_volume_L
            sr_no += 1

            # Sugar / Sorbitol Base (mutually exclusive: EITHER sugar OR sorbitol, NEVER both)
            already_has_sugar = any("sugar" in str(it.get("name", "")).lower() and "sorbitol" not in str(it.get("name", "")).lower() for it in items)
            already_has_sorbitol = any("sorbitol" in str(it.get("name", "")).lower() for it in items)

            # Compute accurate base syrup weight (kg) from bottle fill volume and concentration (% w/v)
            sugar_pct_str = str(sugar_percent or ("30%" if is_sugar_free else "60%")).strip()
            sugar_pct_digits = re.sub(r'[^0-9.]', '', sugar_pct_str)
            try:
                sugar_pct_val = float(sugar_pct_digits) if sugar_pct_digits else (30.0 if is_sugar_free else 60.0)
            except (ValueError, TypeError):
                sugar_pct_val = 30.0 if is_sugar_free else 60.0

            sugar_pct_ratio = sugar_pct_val / 100.0
            base_qty_kg = round(total_volume_L * sugar_pct_ratio, 4)
            base_qty_g = round(base_qty_kg * 1000.0, 1)
            base_qty_mg = round(base_qty_kg * 1_000_000.0, 0)
            sugar_pct_key = f"{int(sugar_pct_val)}%" if sugar_pct_val.is_integer() else f"{sugar_pct_val}%"

            if is_sugar_free:
                # Sugar-free -> only Sorbitol base
                if not already_has_sorbitol and sugar_percent:
                    sorb_rate = self.sorbitol_cost_mapping.get(sugar_pct_key, self.sorbitol_cost_mapping.get(sugar_percent, 216))
                    sorb_cost = sorb_rate * (quantity / 100.0)
                    rate_per_kg = round(sorb_cost / base_qty_kg, 2) if base_qty_kg > 0 else 72.00
                    items.append({
                        "sr_no": sr_no,
                        "group": "B. Base Syrup",
                        "name": f"Sorbitol Base ({sugar_percent} w/v)",
                        "other_name": f"Sorbitol ({sugar_percent})",
                        "type": "Powder",
                        "qty_kg": base_qty_kg,
                        "qty_g": base_qty_g,
                        "qty_mg": base_qty_mg,
                        "qty_per_unit": f"{sugar_percent} w/v",
                        "rate": sorb_rate,
                        "rate_per_kg": rate_per_kg,
                        "cost": round(sorb_cost, 2),
                        "unit_cost_display": f"₹{rate_per_kg:.2f}/kg"
                    })
                    total_other_cost += sorb_cost
                    total_batch_kg += base_qty_kg
                    sr_no += 1
            elif "none" not in st_lower:
                # Standard With Sugar -> only Sugar base
                if not already_has_sugar and sugar_percent:
                    sug_rate = self.sugar_cost_mapping.get(sugar_pct_key, self.sugar_cost_mapping.get(sugar_percent, 168))
                    sug_cost = sug_rate * (quantity / 100.0)
                    rate_per_kg = round(sug_cost / base_qty_kg, 2) if base_qty_kg > 0 else 56.00
                    items.append({
                        "sr_no": sr_no,
                        "group": "B. Base Syrup",
                        "name": f"Sugar Base ({sugar_percent} w/v)",
                        "other_name": f"Sugar ({sugar_percent})",
                        "type": "Powder",
                        "qty_kg": base_qty_kg,
                        "qty_g": base_qty_g,
                        "qty_mg": base_qty_mg,
                        "qty_per_unit": f"{sugar_percent} w/v",
                        "rate": sug_rate,
                        "rate_per_kg": rate_per_kg,
                        "cost": round(sug_cost, 2),
                        "unit_cost_display": f"₹{rate_per_kg:.2f}/kg"
                    })
                    total_other_cost += sug_cost
                    total_batch_kg += base_qty_kg
                    sr_no += 1

        else: # Tablet or Capsule
            scale_factor = quantity / 100.0
            for exc in selected_excipients:
                if not exc.get("selected", True):
                    continue
                name = exc.get("name", "")
                qty_per_100 = float(exc.get("qty_per_100", 0.0))
                rate = float(exc.get("rate", 0.0))
                qty_kg = qty_per_100 * scale_factor
                qty_g = qty_kg * 1000.0
                qty_mg = qty_kg * 1_000_000.0
                cost = qty_kg * rate
                qty_per_unit = f"{qty_mg / quantity:.2f} mg" if quantity > 0 else "0.00 mg"

                items.append({
                    "sr_no": sr_no,
                    "group": exc.get("group", "B. Excipients"),
                    "name": name,
                    "other_name": exc.get("other_name", name),
                    "type": exc.get("type", "Powder"),
                    "qty_kg": round(qty_kg, 6),
                    "qty_g": round(qty_g, 3),
                    "qty_mg": round(qty_mg, 2),
                    "qty_per_unit": qty_per_unit,
                    "rate": round(rate, 2),
                    "cost": round(cost, 2),
                    "unit_cost_display": ""
                })
                total_other_cost += cost
                total_batch_kg += qty_kg
                sr_no += 1

            # Capsule shell cost:
            # - Delayed Release (DR): 0.75
            # - Size 000: Gelatin 0.30, Vegetarian (HPMC) 0.60
            # - Sizes 2, 1, 0, 00: Gelatin 0.18, Vegetarian (HPMC) 0.48
            if product_type == "Capsule":
                c_type_str = str(capsule_type).lower().strip()
                size_str = str(size_or_capsule).strip()

                if "dr" in c_type_str:
                    capsule_unit_cost = 0.75
                    shell_label = "Delayed Release (DR) Shell"
                elif size_str == "000":
                    if "gelatin" in c_type_str or "nonveg" in c_type_str or "non_veg" in c_type_str:
                        capsule_unit_cost = 0.30
                        shell_label = "Size 000 Gelatin (Non-Veg) Shell"
                    else:
                        capsule_unit_cost = 0.60
                        shell_label = "Size 000 Vegetarian (HPMC) Shell"
                elif "gelatin" in c_type_str or "nonveg" in c_type_str or "non_veg" in c_type_str:
                    capsule_unit_cost = 0.18
                    shell_label = "Gelatin (Non-Veg) Shell"
                else:
                    capsule_unit_cost = 0.48
                    shell_label = "Vegetarian (HPMC) Shell"

                capsule_cost = quantity * capsule_unit_cost

                items.append({
                    "sr_no": sr_no,
                    "group": "D. Capsule Shell",
                    "name": f"Empty Capsule Shell (Size {size_or_capsule}, {shell_label})",
                    "other_name": f"{shell_label}",
                    "type": "Shell",
                    "qty_kg": 0.0,
                    "qty_g": 0.0,
                    "qty_mg": 0.0,
                    "qty_per_unit": "1 shell",
                    "rate": capsule_unit_cost,
                    "cost": round(capsule_cost, 2),
                    "unit_cost_display": f"₹{capsule_unit_cost:.2f}/shell"
                })
                total_other_cost += capsule_cost
                sr_no += 1

        # Packaging Calculations
        packaging_summary = self._calculate_packaging(product_type, quantity, size_or_capsule, packaging)

        total_packaging_cost = packaging_summary["total_packaging_cost"]
        total_ingredient_cost = total_active_cost + total_other_cost
        if product_type == "Liquid":
            cfg_liq = master_config_mgr.get_config().get("liquid", {})
            mfg_cfg = cfg_liq.get("manufacturing", {})
            conv_b = float(packaging.get("conversion_rate", mfg_cfg.get("conversion_cost_per_bottle", 1.50)))
            stereo_b = float(packaging.get("stereo_rate", mfg_cfg.get("stereo_cost_per_bottle", 0.05)))
            testing_b = float(packaging.get("testing_rate", mfg_cfg.get("testing_charge_per_bottle", 0.50)))
            margin_b = float(packaging.get("margin_rate", mfg_cfg.get("margin_per_bottle", 1.00)))

            conversion_cost = (conv_b + stereo_b + testing_b) * quantity
            total_cost_with_conversion = total_ingredient_cost + total_packaging_cost + conversion_cost
            profit_margin = margin_b * quantity
            total_batch_cost_with_profit = total_cost_with_conversion + profit_margin
        else:
            conv_b = 0.25
            stereo_b = 0.0
            testing_b = 0.0
            margin_b = 0.0
            conversion_cost = (quantity / 10.0) * 2.5
            total_cost_with_conversion = total_ingredient_cost + total_packaging_cost + conversion_cost
            profit_margin = 0.20 * total_cost_with_conversion
            total_batch_cost_with_profit = total_cost_with_conversion + profit_margin

        # Unit pricing
        profit_per_strip = 0.0
        profit_per_jar = 0.0
        profit_per_loose = 0.0
        rate_per_unit = 0.0
        profit_per_bottle = 0.0
        profit_per_serving = 0.0

        if quantity > 0:
            rate_per_unit = total_batch_cost_with_profit / quantity

        if packaging_summary.get("num_strips", 0) > 0:
            profit_per_strip = total_batch_cost_with_profit / packaging_summary["num_strips"]
        if packaging_summary.get("num_jars", 0) > 0:
            profit_per_jar = total_batch_cost_with_profit / packaging_summary["num_jars"]
        if packaging_summary.get("num_loose", 0) > 0:
            profit_per_loose = total_batch_cost_with_profit / packaging_summary["num_loose"]
        else:
            profit_per_loose = rate_per_unit

        if product_type == "Liquid" and quantity > 0:
            profit_per_bottle = total_batch_cost_with_profit / quantity
            if total_servings > 0:
                profit_per_serving = total_batch_cost_with_profit / total_servings

        # Detailed Commercial Rate Derivation
        cap_per_strip = 10
        strip_size_str = str(packaging.get("strip_size", "1*10"))
        if "*" in strip_size_str:
            try:
                cap_per_strip = int(strip_size_str.split("*")[1])
            except Exception:
                cap_per_strip = 10

        active_cost_per_unit = (total_active_cost / quantity) if quantity > 0 else 0.0
        excipient_cost_per_unit = (total_other_cost / quantity) if quantity > 0 else 0.0
        rm_cost_per_unit = (total_ingredient_cost / quantity) if quantity > 0 else 0.0
        conversion_cost_per_unit = (conversion_cost / quantity) if quantity > 0 else 0.25
        packaging_cost_per_unit = (total_packaging_cost / quantity) if quantity > 0 else 0.0
        margin_per_unit = (profit_margin / quantity) if quantity > 0 else 0.0
        num_strips = packaging_summary.get("num_strips", 0)
        packaging_cost_per_strip = (total_packaging_cost / num_strips) if num_strips > 0 else (packaging_cost_per_unit * cap_per_strip)
        net_mfg_cost_per_unit = (total_cost_with_conversion / quantity) if quantity > 0 else 0.0
        net_mfg_cost_per_strip = (net_mfg_cost_per_unit * cap_per_strip)

        rate_derivation = {
            "batch_quantity": quantity,
            "product_type": product_type,
            "strip_size": strip_size_str,
            "capsules_or_tablets_per_strip": cap_per_strip,
            "num_strips": num_strips,
            "active_ingredients_total": round(total_active_cost, 2),
            "excipients_total": round(total_other_cost, 2),
            "total_raw_material_cost": round(total_ingredient_cost, 2),
            "conversion_charges_total": round(conversion_cost, 2),
            "packaging_material_total": round(total_packaging_cost, 2),
            "net_manufacturing_cost": round(total_cost_with_conversion, 2),
            "profit_margin_20_total": round(profit_margin, 2),
            "total_batch_commercial_val": round(total_batch_cost_with_profit, 2),
            "active_cost_per_unit": round(active_cost_per_unit, 4),
            "excipient_cost_per_unit": round(excipient_cost_per_unit, 4),
            "raw_material_cost_per_unit": round(rm_cost_per_unit, 4),
            "conversion_cost_per_unit": round(conversion_cost_per_unit, 4),
            "packaging_cost_per_unit": round(packaging_cost_per_unit, 4),
            "net_mfg_cost_per_unit": round(net_mfg_cost_per_unit, 4),
            "margin_per_unit": round(margin_per_unit, 4),
            "final_rate_per_unit": round(rate_per_unit, 4),
            "active_cost_per_strip": round(active_cost_per_unit * cap_per_strip, 2),
            "excipient_cost_per_strip": round(excipient_cost_per_unit * cap_per_strip, 2),
            "raw_material_cost_per_strip": round(rm_cost_per_unit * cap_per_strip, 2),
            "conversion_cost_per_strip": round(conversion_cost_per_unit * cap_per_strip, 2),
            "packaging_cost_per_strip": round(packaging_cost_per_strip, 2),
            "net_mfg_cost_per_strip": round(net_mfg_cost_per_strip, 2),
            "margin_per_strip": round(margin_per_unit * cap_per_strip, 2),
            "final_rate_per_strip": round(profit_per_strip, 2),
            "formula_unit": f"Rate Per {product_type} = Total Commercial Value (₹{total_batch_cost_with_profit:,.2f}) ÷ Batch Quantity ({quantity:,}) = ₹{rate_per_unit:.4f}",
            "formula_strip": f"Rate Per Strip = Rate Per {product_type} (₹{rate_per_unit:.4f}) × {cap_per_strip} Units/Strip = ₹{profit_per_strip:.2f}",
            "user_estimated_items": [itm["name"] for itm in items if itm.get("is_user_estimated")]
        }

        liquid_master_card = None
        if product_type == "Liquid":
            sugar_desc = f"{'Sugar Free (Sorbitol)' if is_sugar_free else 'With Sugar'} ({sugar_percent} w/v)" if sugar_type != "None" else "None (Aqueous)"
            table_a_items = []
            for itm in items:
                claim_val = itm.get("qty_per_unit", "-")
                if itm.get("group") != "A. Active Ingredients" and "base" not in itm.get("group", "").lower() and "sugar" not in str(itm.get("name", "")).lower() and "sorbitol" not in str(itm.get("name", "")).lower():
                    claim_val = "-"

                itm_name_l = str(itm.get("name", "")).lower()
                is_base_syrup = "sugar" in itm_name_l or "sorbitol" in itm_name_l
                qty_kg_val = round(float(itm.get("qty_kg", 0.0)), 4)
                if is_base_syrup and qty_kg_val > 0:
                    unit_rate_val = round(float(itm.get("rate_per_kg", itm.get("cost", 0.0) / qty_kg_val)), 2)
                else:
                    unit_rate_val = round(float(itm.get("rate", 0.0)), 2)

                table_a_items.append({
                    "sr_no": itm.get("sr_no"),
                    "group": itm.get("group"),
                    "item_description": itm.get("name"),
                    "claim": claim_val,
                    "std_batch_qty": qty_kg_val,
                    "unit": "LTR" if "water" in itm_name_l else "KG",
                    "unit_rate": unit_rate_val,
                    "total_cost": round(float(itm.get("cost", 0.0)), 2)
                })

            total_vol_l = (quantity * bottle_ml) / 1000.0
            total_pkg_per_btl = packaging_summary.get("cost_per_bottle", 0.0)
            per_bottle_bulk = (total_ingredient_cost / quantity) if quantity > 0 else 0.0
            total_mfg_cost_per_btl = per_bottle_bulk + total_pkg_per_btl + conv_b + stereo_b + testing_b + margin_b

            liquid_master_card = {
                "header": {
                    "company_name": "WALPAR PHARMACEUTICALS",
                    "product_name": payload.get("product_name") or payload.get("brand_name") or f"OSSOFY P suspension {int(bottle_ml)} ml",
                    "batch_size_ltr": round(total_vol_l, 2),
                    "batch_quantity_bottles": quantity,
                    "pack_size": f"{int(bottle_ml) if bottle_ml.is_integer() else bottle_ml} ml",
                    "dosage_basis": f"Each {int(srv_ml) if srv_ml.is_integer() else srv_ml} ml",
                    "sugar_type": sugar_type,
                    "sugar_percent": sugar_percent,
                    "sugar_base_desc": sugar_desc,
                    "yield_percent": float(payload.get("yield_percent", 100.0)),
                    "prepared_by": payload.get("prepared_by", "BMR Production Planning Dept"),
                    "party_name": payload.get("party_name", "Walpar Standard Quotation"),
                    "remarks": payload.get("remarks", "All raw materials & packaging materials as per GMP standards.")
                },
                "table_a_bulk": {
                    "items": table_a_items,
                    "bulk_total_qty_ltr": round(total_vol_l, 2),
                    "bulk_total_cost": round(total_ingredient_cost, 2),
                    "per_bottle_bulk_cost": round(per_bottle_bulk, 2)
                },
                "table_b_semi_finish": {
                    "specs": {
                        "sugar_base": sugar_desc,
                        "bottle": packaging_summary.get("bottle_spec", "WHITE BRUTE BOTTLE"),
                        "ropp": packaging_summary.get("cap_spec", "GOLDEN CAP"),
                        "measuring_cup": f"{int(srv_ml) if srv_ml.is_integer() else srv_ml} ML CUP" if packaging_summary.get("include_measuring_cup") else "N/A",
                        "carton": packaging_summary.get("carton_spec", "uv drip off"),
                        "label": packaging_summary.get("label_spec", "uv varnish roll (120x45)"),
                        "box": packaging_summary.get("box_spec", "60 BOTT"),
                        "dropper": "Standard Dropper" if packaging_summary.get("include_dropper") else "N/A",
                        "insert": "Paper Leaflet" if packaging_summary.get("include_insert") else "N/A"
                    },
                    "costs_per_bottle": {
                        "per_bottle_bulk": round(per_bottle_bulk, 2),
                        "bottle": round(packaging_summary.get("bottle_cost", 0.0), 2),
                        "label": round(packaging_summary.get("label_cost", 0.0), 2),
                        "ropp": round(packaging_summary.get("ropp_cost", 0.0), 2),
                        "carton": round(packaging_summary.get("carton_cost", 0.0), 2),
                        "measuring_cup": round(packaging_summary.get("measuring_cup_cost", 0.0), 2),
                        "box": round(packaging_summary.get("box_cost", 0.0), 2),
                        "dropper": round(packaging_summary.get("dropper_cost", 0.0), 2),
                        "insert": round(packaging_summary.get("insert_cost", 0.0), 2),
                        "pkg_mode": packaging_summary.get("pkg_mode", "MODE_A"),
                        "packaging_total": round(total_pkg_per_btl, 2),
                        "conversion": round(conv_b, 2),
                        "stereo": round(stereo_b, 2),
                        "margin": round(margin_b, 2),
                        "testing_charge": round(testing_b, 2),
                        "total_mfg_cost": round(total_mfg_cost_per_btl, 2),
                        "final_rate": round(total_mfg_cost_per_btl, 2)
                    }
                }
            }

        return {
            "success": True,
            "liquid_master_card": liquid_master_card,
            "product_type": product_type,
            "quantity": quantity,
            "size_or_capsule": size_or_capsule,
            "auto_selected_size": auto_selected_size,
            "total_active_mg": round(total_active_mg, 2),
            "serving_size": serving_size if product_type == "Liquid" else None,
            "total_batch_kg": round(total_batch_kg, 4),
            "items": items,
            "packaging_summary": packaging_summary,
            "rate_derivation": rate_derivation,
            "financials": {
                "active_ingredients_cost": round(total_active_cost, 2),
                "other_ingredients_cost": round(total_other_cost, 2),
                "total_ingredient_cost": round(total_ingredient_cost, 2),
                "total_packaging_cost": round(total_packaging_cost, 2),
                "conversion_cost": round(conversion_cost, 2),
                "total_cost_with_conversion": round(total_cost_with_conversion, 2),
                "profit_margin_20": round(profit_margin, 2),
                "total_batch_cost_with_profit": round(total_batch_cost_with_profit, 2),
                "rate_per_unit": round(rate_per_unit, 4),
                "profit_per_strip": round(profit_per_strip, 2),
                "profit_per_jar": round(profit_per_jar, 2),
                "profit_per_loose": round(profit_per_loose, 2),
                "profit_per_bottle": round(profit_per_bottle, 2),
                "profit_per_serving": round(profit_per_serving, 2),
                "rate_derivation": rate_derivation,
                "missing_rate_count": len(missing_rate_items),
                "missing_rate_items": missing_rate_items,
                "has_missing_rates": len(missing_rate_items) > 0,
                "missing_rate_notice": f"Note: Rate for {', '.join(missing_rate_items)} is currently pending review by Admin. This ingredient rate will be added to this final cost once verified." if missing_rate_items else ""
            },
            "has_missing_rates": len(missing_rate_items) > 0,
            "missing_rate_items": missing_rate_items,
            "missing_rate_notice": f"Note: Rate for {', '.join(missing_rate_items)} is currently pending review by Admin. This ingredient rate will be added to this final cost once verified." if missing_rate_items else ""
        }

    def _calculate_packaging(self, product_type: str, quantity: int, size_or_capsule: Any, packaging: Dict[str, Any]) -> Dict[str, Any]:
        size_str = str(size_or_capsule).replace(" ml", "").replace("ml", "").strip()

        if product_type == "Liquid":
            try:
                b_ml = float(size_str)
            except Exception:
                b_ml = 200.0

            cfg_liq = master_config_mgr.get_config().get("liquid", {})
            std_pkg = cfg_liq.get("standard_packaging", {})
            comps = cfg_liq.get("components", {})

            pkg_mode = str(packaging.get("mode", "MODE_A")).upper()

            # Mode A: Standard Packaging Cost Rule
            if 15.0 <= b_ml <= 60.0:
                std_cost_per_btl = float(std_pkg.get("15_60ml", 8.0))
            elif 60.0 < b_ml <= 100.0:
                std_cost_per_btl = float(std_pkg.get("60_100ml", 10.0))
            elif 100.0 < b_ml <= 200.0:
                std_cost_per_btl = float(std_pkg.get("200ml", 12.0))
            else:
                std_cost_per_btl = float(std_pkg.get("default", 12.0))

            # Mode B: Detailed Itemized Packaging
            bottle_mat = packaging.get("bottle_material", "Plastic")
            bottle_spec = packaging.get("bottle_spec", "WHITE BRUTE BOTTLE" if b_ml >= 200 else "ROUND BRUTE")
            b_key = f"{int(b_ml)}ml"
            default_btl_rate = comps.get("bottles", {}).get(b_key, {}).get("PET", 4.0 if b_ml >= 200 else 2.5)
            bottle_rate = float(packaging.get("bottle_rate", default_btl_rate))

            cap_type = packaging.get("ropp_type", packaging.get("cap_type", "GOLDEN CAP"))
            default_cap_rate = comps.get("caps", {}).get("ROPP Golden", 0.72)
            cap_rate = float(packaging.get("cap_rate", default_cap_rate))

            inc_cup = packaging.get("include_measuring_cup", packaging.get("measuring_cup") != "no")
            cup_rate = float(packaging.get("measuring_cup_rate", 0.30)) if inc_cup else 0.0

            inc_label = packaging.get("include_label", True)
            label_spec = packaging.get("label_type", "uv varnish roll (120x45)")
            label_rate = float(packaging.get("label_rate", 0.65)) if inc_label else 0.0

            inc_carton = packaging.get("include_carton", True)
            carton_spec = packaging.get("carton_type", "uv drip off")
            carton_rate = float(packaging.get("carton_rate", 3.20)) if inc_carton else 0.0

            inc_box = packaging.get("include_box", True)
            box_spec = packaging.get("box_type", "60 BOTT")
            box_rate = float(packaging.get("box_rate", 1.00)) if inc_box else 0.0

            inc_dropper = packaging.get("include_dropper", b_ml < 60.0 and str(packaging.get("dropper_required")).lower() in ["yes", "true"])
            dropper_rate = float(packaging.get("dropper_rate", 1.50)) if inc_dropper else 0.0

            inc_insert = packaging.get("include_insert", str(packaging.get("insert_required")).lower() in ["yes", "true"])
            insert_rate = float(packaging.get("insert_rate", 0.35)) if inc_insert else 0.0

            detailed_cost_per_btl = bottle_rate + cap_rate + cup_rate + label_rate + carton_rate + box_rate + dropper_rate + insert_rate

            if pkg_mode == "MODE_A":
                final_pkg_rate_per_btl = std_cost_per_btl
            else:
                final_pkg_rate_per_btl = detailed_cost_per_btl

            total_packaging_cost = final_pkg_rate_per_btl * quantity

            return {
                "liquid_details": True,
                "pkg_mode": pkg_mode,
                "bottle_cost": round(bottle_rate, 2),
                "ropp_cost": round(cap_rate, 2),
                "measuring_cup_cost": round(cup_rate, 2),
                "carton_cost": round(carton_rate, 2),
                "label_cost": round(label_rate, 2),
                "box_cost": round(box_rate, 2),
                "dropper_cost": round(dropper_rate, 2),
                "insert_cost": round(insert_rate, 2),
                "cost_per_bottle": round(final_pkg_rate_per_btl, 2),
                "total_packaging_cost": round(total_packaging_cost, 2),
                "bottle_spec": bottle_spec,
                "cap_spec": cap_type,
                "label_spec": label_spec,
                "carton_spec": carton_spec,
                "box_spec": box_spec,
                "include_measuring_cup": inc_cup,
                "include_dropper": inc_dropper,
                "include_insert": inc_insert
            }
        else:
            # Tablet / Capsule packaging
            inc_primary = packaging.get("include_primary", True)
            primary_type = packaging.get("primary_type", "STRIP")
            num_jars = 0
            num_strips = 0
            num_loose = 0
            primary_cost = 0.0

            if inc_primary:
                if primary_type == "JAR":
                    jar_type = packaging.get("jar_type", "PET")
                    cap_type = packaging.get("cap_type", "CRC")
                    tab_per_jar = max(1, int(packaging.get("tablets_per_jar", 60)))
                    sticker_laser = packaging.get("sticker_laser", False)

                    num_jars = math.ceil(quantity / tab_per_jar)
                    jar_cost_dict = {"PET": {15: 3.0, 60: 5.3, 180: 6.5, 300: 8.0}, "HDPE": {15: 3.0, 60: 5.5, 180: 7.0, 300: 9.0}}
                    cap_cost_dict = {"CRC": 3.5, "CT": 2.0}
                    sticker_cost_dict = {15: 2.0, 60: 4.5, 180: 4.85, 300: 7.0}

                    jar_c = jar_cost_dict.get(jar_type, {}).get(tab_per_jar, 5.30)
                    cap_c = cap_cost_dict.get(cap_type, 3.50)
                    stick_c = sticker_cost_dict.get(tab_per_jar, 4.50)
                    if sticker_laser:
                        stick_c += 10.00
                    primary_cost = num_jars * (jar_c + cap_c + stick_c)

                elif primary_type == "STRIP":
                    strip_type = packaging.get("strip_type", "Alu Alu")
                    strip_size = packaging.get("strip_size", "1*10")
                    strip_costs = {
                        "Alu Alu": {"1*10": 1.80, "1*15": 2.20, "1*4": 1.00, "1*2": 1.00},
                        "Blister": {"1*10": 0.90, "1*15": 1.80, "1*4": 1.25, "1*2": 0.50},
                        "Pharmafoil": {"1*10": 1.00}
                    }
                    cost_per_strip = strip_costs.get(strip_type, {}).get(strip_size, 1.80)
                    try:
                        cap_per_strip = int(strip_size.split("*")[1]) if "*" in strip_size else 10
                    except Exception:
                        cap_per_strip = 10
                    num_strips = math.ceil(quantity / cap_per_strip) if cap_per_strip > 0 else 1
                    primary_cost = num_strips * cost_per_strip

                else: # Loose
                    loose_type = packaging.get("loose_type", "aluminum pouch")
                    tab_per_loose = max(1, int(packaging.get("tablets_per_loose", 60)))
                    loose_costs = {"aluminum pouch": 15.00, "jar": 30.00, "plastic zip": 9.00}
                    c_loose = loose_costs.get(loose_type, 15.00)
                    num_loose = math.ceil(quantity / tab_per_loose)
                    primary_cost = num_loose * c_loose

            # Secondary Packaging
            inc_secondary = packaging.get("include_secondary", True)
            secondary_cost = 0.0
            if inc_secondary:
                sec_method = packaging.get("secondary_method", "Individual Carton")
                mat_type = packaging.get("material_type", "350 GSM")
                sec_size = packaging.get("secondary_size", "1*10")
                sec_key = f"{sec_method}_{mat_type} {sec_size}"
                secondary_cost_per_unit = {
                    "Individual Carton_350 GSM 1*10": 3.00, "Individual Carton_350 GSM 1*15": 3.00,
                    "Individual Carton_350 GSM 3*10": 4.00, "Individual Carton_350 GSM 6*10": 5.50,
                    "Individual Carton_350 GSM 5*10": 5.50, "Individual Carton_350 GSM 10*10": 6.50,
                    "Individual Carton_300 GSM 1*10": 2.00, "Individual Carton_300 GSM 1*15": 2.00,
                    "Individual Carton_300 GSM 3*10": 3.00, "Individual Carton_300 GSM 6*10": 5.00,
                    "Individual Carton_300 GSM 5*10": 5.00, "Individual Carton_300 GSM 10*10": 5.75,
                    "Box in Box_350 GSM 1*10": 1.40, "Box in Box_350 GSM 3*10": 1.50,
                    "Box in Box_350 GSM 5*10": 1.80, "Box in Box_350 GSM 6*10": 2.15,
                    "Box in Box_300 GSM 1*10": 1.00, "Box in Box_300 GSM 3*10": 1.30,
                    "Box in Box_300 GSM 5*10": 1.50, "Box in Box_300 GSM 6*10": 1.80,
                    "Box in Box_Metallic 1*10": 1.30, "Box in Box_Metallic 3*10": 1.50,
                    "Box in Box_Metallic 5*10": 1.70, "Box in Box_Metallic 6*10": 2.00
                }
                units_needed = num_jars if primary_type == "JAR" else num_strips
                if units_needed == 0:
                    units_needed = max(1, quantity // 10)

                if sec_method == "Box in Box":
                    boxes_needed = math.ceil(units_needed / 10.0)
                    secondary_cost = boxes_needed * secondary_cost_per_unit.get(sec_key, 1.40)
                else:
                    secondary_cost = units_needed * secondary_cost_per_unit.get(sec_key, 3.00)

            # Tertiary Packaging
            inc_tertiary = packaging.get("include_tertiary", True)
            tertiary_cost = 0.0
            if inc_tertiary:
                tert_type = packaging.get("tertiary_type", "5 PLY")
                tert_cost_per_unit = {"5 PLY": 0.30, "7 PLY": 0.40}
                units_needed = num_jars if primary_type == "JAR" else num_strips
                if units_needed == 0:
                    units_needed = max(1, quantity // 10)
                tertiary_cost = units_needed * tert_cost_per_unit.get(tert_type, 0.30)

            # Cellotape & wrapping
            inc_cellotape = packaging.get("include_cellotape", True)
            cellotape_cost = 0.0
            if inc_cellotape:
                units_needed = num_jars if primary_type == "JAR" else num_strips
                if units_needed == 0:
                    units_needed = max(1, quantity // 10)
                cellotape_cost = units_needed * 0.50

            total_packaging_cost = primary_cost + secondary_cost + tertiary_cost + cellotape_cost

            return {
                "liquid_details": False,
                "primary_type": primary_type,
                "num_jars": num_jars,
                "num_strips": num_strips,
                "num_loose": num_loose,
                "primary_packaging_cost": round(primary_cost, 2),
                "secondary_packaging_cost": round(secondary_cost, 2),
                "tertiary_packaging_cost": round(tertiary_cost, 2),
                "cellotape_wrapping_cost": round(cellotape_cost, 2),
                "total_packaging_cost": round(total_packaging_cost, 2)
            }

    def _generate_liquid_excel(self, bmr_data: Dict[str, Any]) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Batch Master Card"

        # Styling
        header_fill = PatternFill(start_color="0A3641", end_color="0A3641", fill_type="solid")
        teal_sub_fill = PatternFill(start_color="134E5E", end_color="134E5E", fill_type="solid")
        group_fill = PatternFill(start_color="E8F4F8", end_color="E8F4F8", fill_type="solid")
        total_fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
        gold_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
        accent_fill = PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid")

        title_font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
        sub_title_font = Font(name="Calibri", size=11, bold=True, color="0A3641")
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        group_font = Font(name="Calibri", size=10, bold=True, color="0A3641")
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)

        thin_side = Side(style='thin', color="CCCCCC")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        l_card = bmr_data.get("liquid_master_card") or {}
        hdr = l_card.get("header", {})
        t_a = l_card.get("table_a_bulk", {})
        t_b = l_card.get("table_b_semi_finish", {})
        specs = t_b.get("specs", {})
        costs = t_b.get("costs_per_bottle", {})

        quantity = int(bmr_data.get("quantity", 5000))
        size_str = str(bmr_data.get("size_or_capsule", "200 ml"))
        try:
            b_ml = float(size_str.replace(" ml", "").replace("ml", "").strip())
        except Exception:
            b_ml = 200.0
        tot_vol_l = float(hdr.get("batch_size_ltr") or ((quantity * b_ml) / 1000.0))
        srv_str = str(bmr_data.get("serving_size", "10 ml"))

        # Row 1: Company Title
        ws.merge_cells("A1:G1")
        c_title = ws["A1"]
        c_title.value = "WALPAR PHARMACEUTICALS - MASTER BATCH CARD"
        c_title.font = title_font
        c_title.fill = header_fill
        c_title.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        # Row 2: Product Name
        ws.merge_cells("A2:G2")
        c_prod = ws["A2"]
        c_prod.value = hdr.get("product_name", f"OSSOFY P suspension {int(b_ml)} ml")
        c_prod.font = sub_title_font
        c_prod.fill = gold_fill
        c_prod.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

        # Row 3: Technical Parameters
        sug_desc = hdr.get("sugar_base_desc") or f"{hdr.get('sugar_type', 'With Sugar')} ({hdr.get('sugar_percent', '60%')} w/v)"
        meta_items = [
            ("A3:A3", f"Batch: {tot_vol_l:.1f}L"),
            ("B3:C3", f"Qty: {quantity:,} Btl"),
            ("D3:D3", f"Pack: {int(b_ml)}ml"),
            ("E3:E3", f"Dose: {hdr.get('dosage_basis', f'Each {srv_str}')}"),
            ("F3:G3", f"Base: {sug_desc}")
        ]
        for rng, txt in meta_items:
            if ":" in rng and rng.split(":")[0] != rng.split(":")[1]:
                ws.merge_cells(rng)
            c = ws[rng.split(":")[0]]
            c.value = txt
            c.font = bold_font
            c.fill = group_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[3].height = 20
        for col_idx in range(1, 8):
            ws.cell(row=3, column=col_idx).border = thin_border

        # Row 4: Section A Title
        ws.merge_cells("A4:G4")
        cA = ws["A4"]
        cA.value = "A. BULK FORMULATION"
        cA.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cA.fill = teal_sub_fill
        cA.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[4].height = 22

        # Row 5: Column Headers
        headers_a = ["Sr No.", "Item Description", "Claim", "Std Batch Qty", "Unit", "Unit Rate (₹)", "Total Cost (₹)"]
        ws.row_dimensions[5].height = 22
        for c_idx, h in enumerate(headers_a, 1):
            cell = ws.cell(row=5, column=c_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Rows 6..N: Items in Table A
        curr_row = 6
        start_bulk_row = curr_row
        bulk_items = t_a.get("items", [])
        if not bulk_items:
            for itm in bmr_data.get("items", []):
                itm_name_l = str(itm.get("name", "")).lower()
                is_base_syrup = "sugar" in itm_name_l or "sorbitol" in itm_name_l
                qty_kg_val = round(float(itm.get("qty_kg", 0.0)), 4)
                if is_base_syrup and qty_kg_val > 0:
                    unit_rate_val = round(float(itm.get("rate_per_kg", itm.get("cost", 0.0) / qty_kg_val)), 2)
                else:
                    unit_rate_val = round(float(itm.get("rate", 0.0)), 2)

                bulk_items.append({
                    "sr_no": itm.get("sr_no"),
                    "item_description": itm.get("name"),
                    "claim": itm.get("qty_per_unit", "-") if itm.get("group") == "A. Active Ingredients" else "-",
                    "std_batch_qty": qty_kg_val,
                    "unit": "LTR" if "water" in itm_name_l else "KG",
                    "unit_rate": unit_rate_val,
                    "total_cost": itm.get("cost", 0.0)
                })

        for itm in bulk_items:
            ws.row_dimensions[curr_row].height = 19
            ws.cell(row=curr_row, column=1, value=itm.get("sr_no")).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=curr_row, column=2, value=itm.get("item_description")).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row=curr_row, column=3, value=itm.get("claim")).alignment = Alignment(horizontal="center", vertical="center")

            c_qty = ws.cell(row=curr_row, column=4, value=float(itm.get("std_batch_qty", 0.0)))
            c_qty.alignment = Alignment(horizontal="right", vertical="center")
            c_qty.number_format = '#,##0.0000'

            ws.cell(row=curr_row, column=5, value=itm.get("unit", "KG")).alignment = Alignment(horizontal="center", vertical="center")

            c_rate = ws.cell(row=curr_row, column=6, value=float(itm.get("unit_rate", 0.0)))
            c_rate.alignment = Alignment(horizontal="right", vertical="center")
            c_rate.number_format = '#,##0.00'

            c_cost = ws.cell(row=curr_row, column=7, value=f"=D{curr_row}*F{curr_row}")
            c_cost.alignment = Alignment(horizontal="right", vertical="center")
            c_cost.number_format = '#,##0.00'

            for c in range(1, 8):
                ws.cell(row=curr_row, column=c).border = thin_border
                ws.cell(row=curr_row, column=c).font = regular_font
            curr_row += 1

        end_bulk_row = curr_row - 1
        bulk_tot_row = curr_row

        # Bulk Total Cost Row
        ws.row_dimensions[bulk_tot_row].height = 22
        ws.merge_cells(start_row=bulk_tot_row, start_column=1, end_row=bulk_tot_row, end_column=3)
        ws.cell(row=bulk_tot_row, column=1, value="BULK TOTAL COST").font = bold_font
        ws.cell(row=bulk_tot_row, column=1).alignment = Alignment(horizontal="right", vertical="center")

        c_tot_qty = ws.cell(row=bulk_tot_row, column=4, value=f"=SUM(D{start_bulk_row}:D{end_bulk_row})")
        c_tot_qty.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_qty.font = bold_font
        c_tot_qty.number_format = '#,##0.0000'

        ws.cell(row=bulk_tot_row, column=5, value="KG").alignment = Alignment(horizontal="center", vertical="center")

        c_tot_cost = ws.cell(row=bulk_tot_row, column=7, value=f"=SUM(G{start_bulk_row}:G{end_bulk_row})")
        c_tot_cost.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_cost.font = Font(name="Calibri", size=11, bold=True, color="0F5132")
        c_tot_cost.number_format = '#,##0.00'

        for c in range(1, 8):
            cell_obj = ws.cell(row=bulk_tot_row, column=c)
            cell_obj.border = thin_border
            cell_obj.fill = total_fill
        curr_row += 1

        # Per Bottle Bulk Row
        per_btl_bulk_row = curr_row
        ws.row_dimensions[per_btl_bulk_row].height = 22
        ws.merge_cells(start_row=per_btl_bulk_row, start_column=1, end_row=per_btl_bulk_row, end_column=5)
        c_p_lbl = ws.cell(row=per_btl_bulk_row, column=1, value="Per bottle bulk:")
        c_p_lbl.font = Font(name="Calibri", size=11, bold=True, color="0A3641")
        c_p_lbl.alignment = Alignment(horizontal="right", vertical="center")

        ws.merge_cells(start_row=per_btl_bulk_row, start_column=6, end_row=per_btl_bulk_row, end_column=7)
        c_p_val = ws.cell(row=per_btl_bulk_row, column=6, value=f"=G{bulk_tot_row}/{quantity}")
        c_p_val.font = Font(name="Calibri", size=11, bold=True, color="0F5132")
        c_p_val.alignment = Alignment(horizontal="right", vertical="center")
        c_p_val.number_format = '"₹"#,##0.00'

        for c in range(1, 8):
            cell_obj = ws.cell(row=per_btl_bulk_row, column=c)
            cell_obj.border = thin_border
            cell_obj.fill = gold_fill
        curr_row += 2

        # Section B: SEMI FINISH
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=7)
        cB = ws.cell(row=curr_row, column=1, value="B. SEMI FINISH & PACKAGING COST SUMMARY")
        cB.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cB.fill = teal_sub_fill
        cB.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[curr_row].height = 22
        curr_row += 1

        # Subheaders for Split Table
        ws.row_dimensions[curr_row].height = 20
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
        c_lh = ws.cell(row=curr_row, column=1, value="Packaging Materials & Specifications")
        c_lh.font = bold_font
        c_lh.fill = accent_fill
        c_lh.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells(start_row=curr_row, start_column=4, end_row=curr_row, end_column=5)
        c_ch = ws.cell(row=curr_row, column=4, value="Cost Component / Element")
        c_ch.font = bold_font
        c_ch.fill = accent_fill
        c_ch.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells(start_row=curr_row, start_column=6, end_row=curr_row, end_column=7)
        c_rh = ws.cell(row=curr_row, column=6, value="Rate / Bottle (₹)")
        c_rh.font = bold_font
        c_rh.fill = accent_fill
        c_rh.alignment = Alignment(horizontal="center", vertical="center")

        for c in range(1, 8):
            ws.cell(row=curr_row, column=c).border = thin_border
        curr_row += 1

        # Packaging spec rows (Left) and Cost items (Right)
        specs_list = [
            ("BOTTLE", specs.get("bottle", "WHITE BRUTE BOTTLE")),
            ("ROPP", specs.get("ropp", "GOLDEN CAP")),
            ("MEASURING CUP", specs.get("measuring_cup", f"{int(b_ml)} ML CUP")),
            ("CARTON", specs.get("carton", "uv drip off")),
            ("LABEL", specs.get("label", "uv varnish roll (120x45)")),
            ("BOX", specs.get("box", "60 BOTT")),
            ("DROPPER", specs.get("dropper", "N/A")),
            ("INSERT", specs.get("insert", "N/A"))
        ]
        active_specs = [(k, v) for k, v in specs_list if v and v != "N/A"]

        cost_breakdown_items = [
            ("Per bottle bulk", f"=F{per_btl_bulk_row}", True),
            ("Bottle", float(costs.get("bottle", 4.00)), False),
            ("Label", float(costs.get("label", 0.65)), False),
            ("ROPP", float(costs.get("ropp", 0.72)), False),
            ("Carton", float(costs.get("carton", 3.20)), False),
            ("Measuring Cup", float(costs.get("measuring_cup", 0.30)), False),
            ("Box", float(costs.get("box", 1.00)), False),
            ("Conversion", float(costs.get("conversion", 1.50)), False),
            ("Stereo", float(costs.get("stereo", 0.05)), False),
            ("Margin", float(costs.get("margin", 1.00)), False),
            ("Testing charge", float(costs.get("testing_charge", 0.50)), False)
        ]

        if float(costs.get("dropper", 0)) > 0:
            cost_breakdown_items.insert(7, ("Dropper", float(costs.get("dropper", 0)), False))
        if float(costs.get("insert", 0)) > 0:
            cost_breakdown_items.insert(8, ("Insert", float(costs.get("insert", 0)), False))

        max_rows = max(len(active_specs), len(cost_breakdown_items))
        start_semi_cost_row = curr_row

        for idx in range(max_rows):
            ws.row_dimensions[curr_row].height = 19
            if idx < len(active_specs):
                sp_k, sp_v = active_specs[idx]
                ws.cell(row=curr_row, column=1, value=sp_k).font = bold_font
                ws.cell(row=curr_row, column=1).alignment = Alignment(horizontal="left", vertical="center")
                ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=3)
                ws.cell(row=curr_row, column=2, value=sp_v).font = regular_font
                ws.cell(row=curr_row, column=2).alignment = Alignment(horizontal="left", vertical="center")
            else:
                ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)

            if idx < len(cost_breakdown_items):
                c_lbl, c_val, is_formula = cost_breakdown_items[idx]
                ws.merge_cells(start_row=curr_row, start_column=4, end_row=curr_row, end_column=5)
                c_l = ws.cell(row=curr_row, column=4, value=c_lbl)
                c_l.font = regular_font
                c_l.alignment = Alignment(horizontal="left", vertical="center")

                ws.merge_cells(start_row=curr_row, start_column=6, end_row=curr_row, end_column=7)
                c_v = ws.cell(row=curr_row, column=6, value=c_val)
                c_v.font = bold_font if is_formula else regular_font
                c_v.alignment = Alignment(horizontal="right", vertical="center")
                c_v.number_format = '#,##0.00'
            else:
                ws.merge_cells(start_row=curr_row, start_column=4, end_row=curr_row, end_column=5)
                ws.merge_cells(start_row=curr_row, start_column=6, end_row=curr_row, end_column=7)

            for c in range(1, 8):
                ws.cell(row=curr_row, column=c).border = thin_border
            curr_row += 1

        end_semi_cost_row = curr_row - 1

        # Total Mfg. Cost Row
        mfg_cost_row = curr_row
        ws.row_dimensions[mfg_cost_row].height = 22
        ws.merge_cells(start_row=mfg_cost_row, start_column=1, end_row=mfg_cost_row, end_column=3)
        ws.merge_cells(start_row=mfg_cost_row, start_column=4, end_row=mfg_cost_row, end_column=5)
        c_m_lbl = ws.cell(row=mfg_cost_row, column=4, value="TOTAL MFG. COST:")
        c_m_lbl.font = Font(name="Calibri", size=11, bold=True, color="0A3641")
        c_m_lbl.alignment = Alignment(horizontal="right", vertical="center")

        ws.merge_cells(start_row=mfg_cost_row, start_column=6, end_row=mfg_cost_row, end_column=7)
        c_m_val = ws.cell(row=mfg_cost_row, column=6, value=f"=SUM(F{start_semi_cost_row}:F{end_semi_cost_row})")
        c_m_val.font = Font(name="Calibri", size=12, bold=True, color="0F5132")
        c_m_val.alignment = Alignment(horizontal="right", vertical="center")
        c_m_val.number_format = '"₹"#,##0.00'

        for c in range(1, 8):
            cell_obj = ws.cell(row=mfg_cost_row, column=c)
            cell_obj.border = thin_border
            cell_obj.fill = total_fill
        curr_row += 1

        # Final Commercial Rate Row
        rate_row = curr_row
        ws.row_dimensions[rate_row].height = 24
        ws.merge_cells(start_row=rate_row, start_column=1, end_row=rate_row, end_column=3)
        ws.merge_cells(start_row=rate_row, start_column=4, end_row=rate_row, end_column=5)
        c_r_lbl = ws.cell(row=rate_row, column=4, value="FINAL RATE / BOTTLE:")
        c_r_lbl.font = Font(name="Calibri", size=11, bold=True, color="0A3641")
        c_r_lbl.alignment = Alignment(horizontal="right", vertical="center")

        ws.merge_cells(start_row=rate_row, start_column=6, end_row=rate_row, end_column=7)
        c_r_val = ws.cell(row=rate_row, column=6, value=f"=F{mfg_cost_row}")
        c_r_val.font = Font(name="Calibri", size=13, bold=True, color="0A3641")
        c_r_val.alignment = Alignment(horizontal="right", vertical="center")
        c_r_val.number_format = '"₹"#,##0.00'

        for c in range(1, 8):
            cell_obj = ws.cell(row=rate_row, column=c)
            cell_obj.border = thin_border
            cell_obj.fill = gold_fill
        curr_row += 2

        # Footer / Commercial Signoff
        signoff_items = [
            ("Party Name:", hdr.get("party_name", "Walpar Standard Quotation")),
            ("Prepared By:", hdr.get("prepared_by", "BMR Production Planning Dept")),
            ("Remarks:", hdr.get("remarks", "All raw materials & packaging materials as per GMP standards."))
        ]
        for s_lbl, s_val in signoff_items:
            ws.row_dimensions[curr_row].height = 18
            ws.cell(row=curr_row, column=1, value=s_lbl).font = bold_font
            ws.cell(row=curr_row, column=1).alignment = Alignment(horizontal="left", vertical="center")
            ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=7)
            c_sv = ws.cell(row=curr_row, column=2, value=s_val)
            c_sv.font = regular_font
            c_sv.alignment = Alignment(horizontal="left", vertical="center")
            for c in range(1, 8):
                ws.cell(row=curr_row, column=c).border = thin_border
            curr_row += 1

        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 36
        ws.column_dimensions['C'].width = 22
        ws.column_dimensions['D'].width = 16
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 16
        ws.column_dimensions['G'].width = 18

        # ── SHEET 2: COMMERCIAL COSTING & RATE DERIVATION ──
        ws2 = wb.create_sheet(title="Costing & Commercial Rates")
        ws2.merge_cells("A1:H1")
        ws2["A1"] = "WALPAR PHARMACEUTICALS - COMMERCIAL COSTING & RATE DERIVATION (LIQUID ORAL)"
        ws2["A1"].font = title_font
        ws2["A1"].fill = header_fill
        ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[1].height = 28

        ws2.merge_cells("A2:H2")
        ws2["A2"] = "Dynamic Excel Model | Modifying ingredient rates instantly recalculates rate per bottle and commercial batch value"
        ws2["A2"].font = Font(name="Calibri", size=10, italic=True, color="444444")
        ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[2].height = 20

        ws2.row_dimensions[3].height = 24
        driver_items = [
            (1, "Batch Quantity (Bottles):", 2, quantity, '#,##0'),
            (3, "Bottle Fill Volume:", 4, f"{int(b_ml)} ml", '@'),
            (5, "Dosage Basis:", 6, hdr.get('dosage_basis', f'Each {srv_str}'), '@'),
            (7, "Standard Batch Ltr:", 8, tot_vol_l, '#,##0.00')
        ]
        for l_col, label, v_col, val, fmt in driver_items:
            c_lbl = ws2.cell(row=3, column=l_col, value=label)
            c_lbl.font = bold_font
            c_lbl.fill = accent_fill
            c_lbl.alignment = Alignment(horizontal="right", vertical="center")
            c_lbl.border = thin_border

            c_val = ws2.cell(row=3, column=v_col, value=val)
            c_val.font = Font(name="Calibri", size=11, bold=True, color="0A3641")
            c_val.fill = gold_fill
            c_val.alignment = Alignment(horizontal="center", vertical="center")
            c_val.border = thin_border
            if fmt != '@':
                c_val.number_format = fmt

        ws2.cell(row=5, column=1, value="1. ACTIVE RAW MATERIALS COSTING").font = group_font
        headers2 = ["Sr", "Active Ingredient Name", "Std Qty (Kg)", "Rate (₹ / Kg)", "Total Cost (₹)", "Cost / Bottle (₹)", "Rate Source"]
        ws2.row_dimensions[6].height = 24
        for c_idx, h in enumerate(headers2, 1):
            c = ws2.cell(row=6, column=c_idx, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border

        r_idx = 7
        start_act_row = r_idx
        for itm in bmr_data.get("items", []):
            if itm.get("group") == "A. Active Ingredients":
                src = "User Estimated" if itm.get("is_user_estimated") else ("Official DB" if itm.get("is_rate_available") else "Pending Admin")
                qty_val = float(itm.get("qty_kg", 0.0))
                rate_val = float(itm.get("rate", 0.0))

                ws2.row_dimensions[r_idx].height = 20
                ws2.cell(row=r_idx, column=1, value=itm.get("sr_no")).alignment = Alignment(horizontal="center", vertical="center")
                ws2.cell(row=r_idx, column=2, value=itm.get("name")).alignment = Alignment(horizontal="left", vertical="center")

                c3 = ws2.cell(row=r_idx, column=3, value=qty_val)
                c3.alignment = Alignment(horizontal="right", vertical="center")
                c3.number_format = '#,##0.0000'

                c4 = ws2.cell(row=r_idx, column=4, value=rate_val)
                c4.alignment = Alignment(horizontal="right", vertical="center")
                c4.number_format = '#,##0.00'
                c4.font = bold_font

                c5 = ws2.cell(row=r_idx, column=5, value=f"=C{r_idx}*D{r_idx}")
                c5.alignment = Alignment(horizontal="right", vertical="center")
                c5.number_format = '#,##0.00'

                c6 = ws2.cell(row=r_idx, column=6, value=f"=E{r_idx}/$B$3")
                c6.alignment = Alignment(horizontal="right", vertical="center")
                c6.number_format = '#,##0.0000'

                ws2.cell(row=r_idx, column=7, value=src).alignment = Alignment(horizontal="center", vertical="center")

                for c in range(1, 8):
                    ws2.cell(row=r_idx, column=c).border = thin_border
                    if c != 4:
                        ws2.cell(row=r_idx, column=c).font = regular_font
                r_idx += 1

        end_act_row = r_idx - 1
        tot_act_row = r_idx

        ws2.row_dimensions[tot_act_row].height = 22
        ws2.merge_cells(start_row=tot_act_row, start_column=1, end_row=tot_act_row, end_column=4)
        c_tot = ws2.cell(row=tot_act_row, column=1, value="TOTAL ACTIVE RAW MATERIALS COST")
        c_tot.font = bold_font
        c_tot.alignment = Alignment(horizontal="right", vertical="center")

        if end_act_row >= start_act_row:
            c_tot_v = ws2.cell(row=tot_act_row, column=5, value=f"=SUM(E{start_act_row}:E{end_act_row})")
        else:
            c_tot_v = ws2.cell(row=tot_act_row, column=5, value=0.0)
        c_tot_v.font = bold_font
        c_tot_v.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_v.number_format = '#,##0.00'

        c_tot_u = ws2.cell(row=tot_act_row, column=6, value=f"=E{tot_act_row}/$B$3")
        c_tot_u.font = bold_font
        c_tot_u.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_u.number_format = '#,##0.0000'

        for c in range(1, 8):
            cell_obj = ws2.cell(row=tot_act_row, column=c)
            cell_obj.border = thin_border
            cell_obj.fill = group_fill

        r_idx += 2

        # Section 1B: Excipients
        ws2.cell(row=r_idx, column=1, value="1B. INACTIVE EXCIPIENTS, FLAVOURS & BASES").font = group_font
        r_idx += 1
        headers_exc = ["Sr", "Excipient / Material Name", "Qty (Kg / Ltr)", "Rate (₹)", "Total Cost (₹)", "Cost / Bottle (₹)", "Category"]
        ws2.row_dimensions[r_idx].height = 24
        for c_idx, h in enumerate(headers_exc, 1):
            c = ws2.cell(row=r_idx, column=c_idx, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border

        r_idx += 1
        start_exc_row = r_idx
        for itm in bmr_data.get("items", []):
            if itm.get("group") != "A. Active Ingredients":
                qty_val = float(itm.get("qty_kg", 0.0))
                rate_val = float(itm.get("rate", 0.0))

                ws2.row_dimensions[r_idx].height = 19
                ws2.cell(row=r_idx, column=1, value=itm.get("sr_no")).alignment = Alignment(horizontal="center", vertical="center")
                ws2.cell(row=r_idx, column=2, value=itm.get("name")).alignment = Alignment(horizontal="left", vertical="center")

                c3 = ws2.cell(row=r_idx, column=3, value=qty_val)
                c3.alignment = Alignment(horizontal="right", vertical="center")
                c3.number_format = '#,##0.0000'

                c4 = ws2.cell(row=r_idx, column=4, value=rate_val)
                c4.alignment = Alignment(horizontal="right", vertical="center")
                c4.number_format = '#,##0.00'

                c5 = ws2.cell(row=r_idx, column=5, value=f"=C{r_idx}*D{r_idx}")
                c5.alignment = Alignment(horizontal="right", vertical="center")
                c5.number_format = '#,##0.00'

                c6 = ws2.cell(row=r_idx, column=6, value=f"=E{r_idx}/$B$3")
                c6.alignment = Alignment(horizontal="right", vertical="center")
                c6.number_format = '#,##0.0000'

                ws2.cell(row=r_idx, column=7, value=itm.get("group", "")).alignment = Alignment(horizontal="center", vertical="center")

                for c in range(1, 8):
                    ws2.cell(row=r_idx, column=c).border = thin_border
                    ws2.cell(row=r_idx, column=c).font = regular_font
                r_idx += 1

        end_exc_row = r_idx - 1
        tot_exc_row = r_idx

        ws2.row_dimensions[tot_exc_row].height = 22
        ws2.merge_cells(start_row=tot_exc_row, start_column=1, end_row=tot_exc_row, end_column=4)
        c_tot_e = ws2.cell(row=tot_exc_row, column=1, value="TOTAL INACTIVE EXCIPIENTS & BASES COST")
        c_tot_e.font = bold_font
        c_tot_e.alignment = Alignment(horizontal="right", vertical="center")

        if end_exc_row >= start_exc_row:
            c_tot_ev = ws2.cell(row=tot_exc_row, column=5, value=f"=SUM(E{start_exc_row}:E{end_exc_row})")
        else:
            c_tot_ev = ws2.cell(row=tot_exc_row, column=5, value=0.0)
        c_tot_ev.font = bold_font
        c_tot_ev.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_ev.number_format = '#,##0.00'

        c_tot_eu = ws2.cell(row=tot_exc_row, column=6, value=f"=E{tot_exc_row}/$B$3")
        c_tot_eu.font = bold_font
        c_tot_eu.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_eu.number_format = '#,##0.0000'

        for c in range(1, 8):
            cell_obj = ws2.cell(row=tot_exc_row, column=c)
            cell_obj.border = thin_border
            cell_obj.fill = group_fill

        r_idx += 2

        # Section 2: Derivation Table
        ws2.cell(row=r_idx, column=1, value="2. LIQUID ORAL COMMERCIAL COST DERIVATION & RATE PER BOTTLE").font = group_font
        r_idx += 1
        deriv_headers = ["Cost Component / Element", "Basis of Calculation", "Total Amount (₹)", "Cost / Bottle (₹)"]
        ws2.row_dimensions[r_idx].height = 24
        for c_idx, h in enumerate(deriv_headers, 1):
            c = ws2.cell(row=r_idx, column=c_idx, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border

        r_idx += 1
        r_d_act = r_idx
        ws2.row_dimensions[r_idx].height = 20
        ws2.cell(row=r_idx, column=1, value="A. Active Raw Materials Cost").font = regular_font
        ws2.cell(row=r_idx, column=2, value="Sum of active ingredients in formulation").font = regular_font
        ws2.cell(row=r_idx, column=3, value=f"=E{tot_act_row}").number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=4, value=f"=C{r_idx}/$B$3").number_format = '#,##0.00'
        for c in range(1, 5):
            ws2.cell(row=r_idx, column=c).border = thin_border
        r_idx += 1

        r_d_exc = r_idx
        ws2.row_dimensions[r_idx].height = 20
        ws2.cell(row=r_idx, column=1, value="B. Inactive Excipients & Bases Cost").font = regular_font
        ws2.cell(row=r_idx, column=2, value="Sugar/Sorbitol, Glycerin, Flavours, Xanthan, Water").font = regular_font
        ws2.cell(row=r_idx, column=3, value=f"=E{tot_exc_row}").number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=4, value=f"=C{r_idx}/$B$3").number_format = '#,##0.00'
        for c in range(1, 5):
            ws2.cell(row=r_idx, column=c).border = thin_border
        r_idx += 1

        r_d_pkg = r_idx
        pkg_cost_per_btl = float(costs.get("packaging_total", 9.87))
        ws2.row_dimensions[r_idx].height = 20
        ws2.cell(row=r_idx, column=1, value="C. Complete Packaging Cost").font = regular_font
        ws2.cell(row=r_idx, column=2, value="Bottle, Cap, Measuring Cup, Carton, Label, Box").font = regular_font
        ws2.cell(row=r_idx, column=4, value=pkg_cost_per_btl).number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=3, value=f"=D{r_idx}*$B$3").number_format = '#,##0.00'
        for c in range(1, 5):
            ws2.cell(row=r_idx, column=c).border = thin_border
        r_idx += 1

        r_d_conv = r_idx
        conv_stereo_per_btl = float(costs.get("conversion", 1.50)) + float(costs.get("stereo", 0.05))
        ws2.row_dimensions[r_idx].height = 20
        ws2.cell(row=r_idx, column=1, value="D. Conversion & Stereo Charges").font = regular_font
        ws2.cell(row=r_idx, column=2, value="Standard liquid manufacturing & processing conversion").font = regular_font
        ws2.cell(row=r_idx, column=4, value=conv_stereo_per_btl).number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=3, value=f"=D{r_idx}*$B$3").number_format = '#,##0.00'
        for c in range(1, 5):
            ws2.cell(row=r_idx, column=c).border = thin_border
        r_idx += 1

        r_d_mrg = r_idx
        mrg_test_per_btl = float(costs.get("margin", 1.00)) + float(costs.get("testing_charge", 0.50))
        ws2.row_dimensions[r_idx].height = 20
        ws2.cell(row=r_idx, column=1, value="E. Testing & Operating Margin").font = regular_font
        ws2.cell(row=r_idx, column=2, value="Quality testing charges & operating margin").font = regular_font
        ws2.cell(row=r_idx, column=4, value=mrg_test_per_btl).number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=3, value=f"=D{r_idx}*$B$3").number_format = '#,##0.00'
        for c in range(1, 5):
            ws2.cell(row=r_idx, column=c).border = thin_border
        r_idx += 1

        r_d_tot = r_idx
        ws2.row_dimensions[r_idx].height = 24
        ws2.cell(row=r_idx, column=1, value="F. TOTAL COMMERCIAL VALUATION").font = bold_font
        ws2.cell(row=r_idx, column=2, value=f"Final commercial quotation for {quantity:,} bottles").font = regular_font
        ws2.cell(row=r_idx, column=3, value=f"=SUM(C{r_d_act}:C{r_d_mrg})").number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=4, value=f"=C{r_idx}/$B$3").number_format = '#,##0.00'
        for c in range(1, 5):
            ws2.cell(row=r_idx, column=c).border = thin_border
            ws2.cell(row=r_idx, column=c).fill = total_fill
            if c >= 3:
                ws2.cell(row=r_idx, column=c).font = bold_font
        r_idx += 2

        # Section 3: Official KPI Summary
        ws2.row_dimensions[r_idx].height = 26
        ws2.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=2)
        c_k_lbl = ws2.cell(row=r_idx, column=1, value="* FINAL RATE PER BOTTLE:")
        c_k_lbl.font = Font(name="Calibri", size=12, bold=True, color="0A3641")
        c_k_lbl.alignment = Alignment(horizontal="left", vertical="center")

        ws2.merge_cells(start_row=r_idx, start_column=3, end_row=r_idx, end_column=4)
        c_k_val = ws2.cell(row=r_idx, column=3, value=f"=D{r_d_tot}")
        c_k_val.font = Font(name="Calibri", size=14, bold=True, color="0F5132")
        c_k_val.alignment = Alignment(horizontal="center", vertical="center")
        c_k_val.number_format = '"₹"#,##0.00'
        c_k_val.fill = total_fill
        for c in range(1, 5):
            ws2.cell(row=r_idx, column=c).border = thin_border
        r_idx += 1

        ws2.row_dimensions[r_idx].height = 26
        ws2.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=2)
        c_k_tot_lbl = ws2.cell(row=r_idx, column=1, value=f"* TOTAL COMMERCIAL BATCH VALUE ({quantity:,} BOTTLES):")
        c_k_tot_lbl.font = Font(name="Calibri", size=12, bold=True, color="0A3641")
        c_k_tot_lbl.alignment = Alignment(horizontal="left", vertical="center")

        ws2.merge_cells(start_row=r_idx, start_column=3, end_row=r_idx, end_column=4)
        c_k_tot_val = ws2.cell(row=r_idx, column=3, value=f"=C{r_d_tot}")
        c_k_tot_val.font = Font(name="Calibri", size=14, bold=True, color="0A3641")
        c_k_tot_val.alignment = Alignment(horizontal="center", vertical="center")
        c_k_tot_val.number_format = '"₹"#,##0.00'
        c_k_tot_val.fill = gold_fill
        for c in range(1, 5):
            ws2.cell(row=r_idx, column=c).border = thin_border

        for col in ws2.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and not cell.coordinate in ws2.merged_cells:
                    max_len = len(val_str)
            ws2.column_dimensions[col_letter].width = max(max_len + 3, 15)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def generate_excel(self, bmr_data: Dict[str, Any]) -> io.BytesIO:
        """
        Generates a beautifully formatted openpyxl workbook with live Excel formulas.
        Modifying any ingredient rate automatically recalculates total batch value,
        rate per unit (tablet/capsule), and rate per strip / pack in real time.
        """
        product_type = bmr_data.get("product_type", "Tablet")
        if product_type == "Liquid":
            return self._generate_liquid_excel(bmr_data)

        wb = Workbook()
        ws = wb.active
        ws.title = "Batch Master Card"

        # Theme Colors: Professional Walpar Corporate Teal & Deep Blue
        header_fill = PatternFill(start_color="0A3641", end_color="0A3641", fill_type="solid")
        group_fill = PatternFill(start_color="E8F4F8", end_color="E8F4F8", fill_type="solid")
        total_fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
        summary_fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
        gold_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
        rate_fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
        accent_fill = PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid")

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        group_font = Font(name="Calibri", size=11, bold=True, color="0A3641")
        total_font = Font(name="Calibri", size=11, bold=True, color="0F5132")
        title_font = Font(name="Calibri", size=15, bold=True, color="0A3641")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)
        kpi_title_font = Font(name="Calibri", size=11, bold=True, color="0A3641")
        kpi_val_font = Font(name="Calibri", size=13, bold=True, color="0F5132")
        kpi_tot_font = Font(name="Calibri", size=13, bold=True, color="0A3641")

        thin_side = Side(style='thin', color="CCCCCC")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        product_type = bmr_data.get("product_type", "Tablet")
        quantity = int(bmr_data.get("quantity", 100000))
        size_or_capsule = bmr_data.get("size_or_capsule", "")
        serving_size = bmr_data.get("serving_size", "")
        pkg = bmr_data.get("packaging", {}) or bmr_data.get("packaging_summary", {})
        prim_type = pkg.get("primary_type", "STRIP")
        fin = bmr_data.get("financials", {})
        rd = bmr_data.get("rate_derivation", {}) or fin.get("rate_derivation", {})

        # Packaging parameters
        if prim_type == "STRIP":
            cap_per_strip = rd.get("capsules_or_tablets_per_strip", 10)
            pack_col_name = "Cost / Strip (₹)"
            pack_label = "STRIP"
            pack_size_val = cap_per_strip
        elif prim_type == "JAR":
            tab_per_jar = max(1, int(pkg.get("tablets_per_jar", 60)))
            pack_col_name = "Cost / Jar (₹)"
            pack_label = "JAR"
            pack_size_val = tab_per_jar
        elif product_type == "Liquid":
            pack_col_name = "Cost / Bottle (₹)"
            pack_label = "BOTTLE"
            pack_size_val = 1
        else: # Loose
            tab_per_loose = max(1, int(pkg.get("tablets_per_loose", 60)))
            pack_col_name = "Cost / Pack (₹)"
            pack_label = "PACK"
            pack_size_val = tab_per_loose

        # ═════════════════════════════════════════════════════════════════
        # SHEET 1: BATCH MASTER CARD (TECHNICAL BMR FORMULATION)
        # ═════════════════════════════════════════════════════════════════
        ws.merge_cells("A1:K1")
        title_cell = ws["A1"]
        title_cell.value = "WALPAR PHARMACEUTICALS - BATCH MASTER CARD (BMR)"
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        info_text = f"Product: {product_type} | Batch Size: {quantity:,} units | Dosage/Size: {size_or_capsule}"
        if serving_size:
            info_text += f" | Serving: {serving_size}"

        ws.merge_cells("A2:K2")
        info_cell = ws["A2"]
        info_cell.value = info_text
        info_cell.font = bold_font
        info_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        headers = [
            "Sr No", "Group / Stage", "Ingredient Name", "Other / Chemical Name",
            "Type", "Qty (Kg)", "Qty (g)", "Qty (mg)", "Qty per Unit"
        ]
        ws.row_dimensions[4].height = 24
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        current_row = 5
        start_bmr_item_row = current_row
        last_group = None

        for item in bmr_data.get("items", []):
            group = item.get("group", "")
            if group != last_group:
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
                g_cell = ws.cell(row=current_row, column=1, value=f"  {group.upper()}")
                g_cell.font = group_font
                g_cell.fill = group_fill
                g_cell.alignment = Alignment(horizontal="left", vertical="center")
                for c in range(1, 10):
                    ws.cell(row=current_row, column=c).border = thin_border
                ws.row_dimensions[current_row].height = 20
                current_row += 1
                last_group = group

            ws.row_dimensions[current_row].height = 19
            ws.cell(row=current_row, column=1, value=item.get("sr_no")).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=current_row, column=2, value=item.get("group")).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row=current_row, column=3, value=item.get("name")).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row=current_row, column=4, value=item.get("other_name")).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row=current_row, column=5, value=item.get("type")).alignment = Alignment(horizontal="center", vertical="center")

            # Qty (Kg)
            c_kg = ws.cell(row=current_row, column=6, value=float(item.get("qty_kg", 0.0)))
            c_kg.alignment = Alignment(horizontal="right", vertical="center")
            c_kg.number_format = '#,##0.0000'

            # Qty (g) = Qty (Kg) * 1000 (Formula)
            c_g = ws.cell(row=current_row, column=7, value=f"=F{current_row}*1000")
            c_g.alignment = Alignment(horizontal="right", vertical="center")
            c_g.number_format = '#,##0.0'

            # Qty (mg) = Qty (Kg) * 1,000,000 (Formula)
            c_mg = ws.cell(row=current_row, column=8, value=f"=F{current_row}*1000000")
            c_mg.alignment = Alignment(horizontal="right", vertical="center")
            c_mg.number_format = '#,##0'

            ws.cell(row=current_row, column=9, value=item.get("qty_per_unit")).alignment = Alignment(horizontal="left", vertical="center")

            for c in range(1, 10):
                ws.cell(row=current_row, column=c).font = regular_font
                ws.cell(row=current_row, column=c).border = thin_border

            current_row += 1

        end_bmr_item_row = current_row - 1

        # Total Row in Sheet 1
        ws.row_dimensions[current_row].height = 22
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        tot_lbl = ws.cell(row=current_row, column=1, value="TOTAL FORMULATION BATCH WEIGHT")
        tot_lbl.font = total_font
        tot_lbl.fill = total_fill
        tot_lbl.alignment = Alignment(horizontal="center", vertical="center")
        for c in range(1, 6):
            ws.cell(row=current_row, column=c).border = thin_border

        cell_kg = ws.cell(row=current_row, column=6, value=f"=SUM(F{start_bmr_item_row}:F{end_bmr_item_row})")
        cell_kg.font = total_font
        cell_kg.fill = total_fill
        cell_kg.alignment = Alignment(horizontal="right", vertical="center")
        cell_kg.number_format = '#,##0.0000'
        cell_kg.border = thin_border

        cell_g = ws.cell(row=current_row, column=7, value=f"=F{current_row}*1000")
        cell_g.font = total_font
        cell_g.fill = total_fill
        cell_g.alignment = Alignment(horizontal="right", vertical="center")
        cell_g.number_format = '#,##0.0'
        cell_g.border = thin_border

        cell_mg = ws.cell(row=current_row, column=8, value=f"=F{current_row}*1000000")
        cell_mg.font = total_font
        cell_mg.fill = total_fill
        cell_mg.alignment = Alignment(horizontal="right", vertical="center")
        cell_mg.number_format = '#,##0'
        cell_mg.border = thin_border

        cell_comp = ws.cell(row=current_row, column=9, value="100% Master Formula")
        cell_comp.font = total_font
        cell_comp.fill = total_fill
        cell_comp.alignment = Alignment(horizontal="center", vertical="center")
        cell_comp.border = thin_border

        total_bmr_row = current_row
        current_row += 2

        # Technical Master Summary Section
        summary_rows = [
            ("Total Formulation Batch Weight (Kg):", f"=TEXT(F{total_bmr_row}, \"0.0000\") & \" Kg\""),
            ("Total Formulation Batch Weight (g):", f"=TEXT(G{total_bmr_row}, \"#,##0.0\") & \" g\""),
            ("Total Batch Pack Quantity:", f"{quantity:,} Units"),
            ("Product Classification / Type:", f"{product_type}"),
            ("Primary Packaging System:", f"{pkg.get('primary_type', 'Standard')} ({pkg.get('strip_size') or pkg.get('tablets_per_jar', '')})".strip()),
            ("Secondary Packaging System:", f"{pkg.get('secondary_method', 'Carton')} - {pkg.get('material_type') or pkg.get('carton_material', 'Standard')}"),
            ("Quality & Regulatory Standard:", "Conforms to In-house / FSSAI / GMP Standards")
        ]

        for label, val in summary_rows:
            ws.merge_cells(start_row=current_row, start_column=5, end_row=current_row, end_column=7)
            lbl_c = ws.cell(row=current_row, column=5, value=label)
            lbl_c.font = bold_font if "Total" in label else regular_font
            lbl_c.alignment = Alignment(horizontal="right", vertical="center")
            lbl_c.fill = summary_fill

            ws.merge_cells(start_row=current_row, start_column=8, end_row=current_row, end_column=9)
            val_c = ws.cell(row=current_row, column=8, value=val)
            val_c.font = bold_font
            val_c.alignment = Alignment(horizontal="left", vertical="center")
            val_c.fill = summary_fill

            for c in range(5, 10):
                ws.cell(row=current_row, column=c).border = thin_border
            current_row += 1

        # Adjust column widths for Sheet 1
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and not cell.coordinate in ws.merged_cells:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # ═════════════════════════════════════════════════════════════════
        # SHEET 2: COMMERCIAL COSTING & RATE DERIVATION (DYNAMIC FORMULAS)
        # ═════════════════════════════════════════════════════════════════
        ws2 = wb.create_sheet(title="Costing & Commercial Rates")

        # Row 1: Title
        ws2.merge_cells("A1:H1")
        ws2["A1"] = "WALPAR PHARMACEUTICALS - COMMERCIAL COSTING & RATE DERIVATION"
        ws2["A1"].font = title_font
        ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[1].height = 28

        # Row 2: Subtitle
        ws2.merge_cells("A2:H2")
        ws2["A2"] = "Dynamic Excel Model | Modifying ingredient rates instantly recalculates rate per strip and unit in real time"
        ws2["A2"].font = Font(name="Calibri", size=10, italic=True, color="444444")
        ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[2].height = 20

        # Row 3: Parameter Drivers
        ws2.row_dimensions[3].height = 24
        driver_items = [
            (1, "Batch Quantity (Units):", 2, quantity, '#,##0'),
            (3, f"Units / {pack_label}:", 4, pack_size_val, '0'),
            (5, "Product Type:", 6, product_type, '@'),
            (7, "Packaging System:", 8, f"{prim_type} ({rd.get('strip_size') or pkg.get('strip_size') or pack_size_val})", '@')
        ]
        for l_col, label, v_col, val, fmt in driver_items:
            c_lbl = ws2.cell(row=3, column=l_col, value=label)
            c_lbl.font = bold_font
            c_lbl.fill = accent_fill
            c_lbl.alignment = Alignment(horizontal="right", vertical="center")
            c_lbl.border = thin_border

            c_val = ws2.cell(row=3, column=v_col, value=val)
            c_val.font = Font(name="Calibri", size=11, bold=True, color="0A3641")
            c_val.fill = gold_fill
            c_val.alignment = Alignment(horizontal="center", vertical="center")
            c_val.border = thin_border
            if fmt != '@':
                c_val.number_format = fmt

        # Section 1: Active Raw Materials Cost
        ws2.cell(row=5, column=1, value="1. ACTIVE RAW MATERIALS COSTING").font = group_font
        headers2 = ["Sr", "Active Ingredient Name", "Qty (Kg)", "Rate (₹ / Kg)", "Total Cost (₹)", "Cost / Unit (₹)", "Rate Source"]
        ws2.row_dimensions[6].height = 24
        for c_idx, h in enumerate(headers2, 1):
            c = ws2.cell(row=6, column=c_idx, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border

        r_idx = 7
        start_act_row = r_idx

        for itm in bmr_data.get("items", []):
            if itm.get("group") == "A. Active Ingredients":
                src = "User Estimated" if itm.get("is_user_estimated") else ("Official DB" if itm.get("is_rate_available") else "Pending Admin")
                qty_val = float(itm.get("qty_kg", 0.0))
                rate_val = float(itm.get("rate", 0.0))

                ws2.row_dimensions[r_idx].height = 20
                c1 = ws2.cell(row=r_idx, column=1, value=itm.get("sr_no"))
                c1.alignment = Alignment(horizontal="center", vertical="center")

                c2 = ws2.cell(row=r_idx, column=2, value=itm.get("name"))
                c2.alignment = Alignment(horizontal="left", vertical="center")

                c3 = ws2.cell(row=r_idx, column=3, value=qty_val)
                c3.alignment = Alignment(horizontal="right", vertical="center")
                c3.number_format = '#,##0.0000'

                c4 = ws2.cell(row=r_idx, column=4, value=rate_val)
                c4.alignment = Alignment(horizontal="right", vertical="center")
                c4.number_format = '#,##0.00'
                c4.font = bold_font

                c5 = ws2.cell(row=r_idx, column=5, value=f"=C{r_idx}*D{r_idx}")
                c5.alignment = Alignment(horizontal="right", vertical="center")
                c5.number_format = '#,##0.00'

                c6 = ws2.cell(row=r_idx, column=6, value=f"=E{r_idx}/$B$3")
                c6.alignment = Alignment(horizontal="right", vertical="center")
                c6.number_format = '#,##0.0000'

                c7 = ws2.cell(row=r_idx, column=7, value=src)
                c7.alignment = Alignment(horizontal="center", vertical="center")

                for c in range(1, 8):
                    ws2.cell(row=r_idx, column=c).border = thin_border
                    if c != 4:
                        ws2.cell(row=r_idx, column=c).font = regular_font
                r_idx += 1

        end_act_row = r_idx - 1
        tot_act_row = r_idx

        # Total Active Cost Row
        ws2.row_dimensions[tot_act_row].height = 22
        ws2.merge_cells(start_row=tot_act_row, start_column=1, end_row=tot_act_row, end_column=4)
        c_tot = ws2.cell(row=tot_act_row, column=1, value="TOTAL ACTIVE RAW MATERIALS COST")
        c_tot.font = bold_font
        c_tot.alignment = Alignment(horizontal="right", vertical="center")

        if end_act_row >= start_act_row:
            c_tot_v = ws2.cell(row=tot_act_row, column=5, value=f"=SUM(E{start_act_row}:E{end_act_row})")
        else:
            c_tot_v = ws2.cell(row=tot_act_row, column=5, value=0.0)
        c_tot_v.font = bold_font
        c_tot_v.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_v.number_format = '#,##0.00'

        c_tot_u = ws2.cell(row=tot_act_row, column=6, value=f"=E{tot_act_row}/$B$3")
        c_tot_u.font = bold_font
        c_tot_u.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_u.number_format = '#,##0.0000'

        for c in range(1, 8):
            cell_obj = ws2.cell(row=tot_act_row, column=c)
            cell_obj.border = thin_border
            cell_obj.fill = group_fill

        r_idx += 2

        # Section 1B: Inactive Excipients, Binders & Shell Costing
        ws2.cell(row=r_idx, column=1, value="1B. INACTIVE EXCIPIENTS, BINDERS & SHELL COSTING").font = group_font
        r_idx += 1
        headers_exc = ["Sr", "Excipient / Material Name", "Qty (Kg / Units)", "Rate (₹)", "Total Cost (₹)", "Cost / Unit (₹)", "Category / Stage"]
        ws2.row_dimensions[r_idx].height = 24
        for c_idx, h in enumerate(headers_exc, 1):
            c = ws2.cell(row=r_idx, column=c_idx, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border

        r_idx += 1
        start_exc_row = r_idx

        for itm in bmr_data.get("items", []):
            if itm.get("group") != "A. Active Ingredients":
                qty_val = float(itm.get("qty_kg", 0.0))
                rate_val = float(itm.get("rate", 0.0))
                if itm.get("type") == "Shell" and qty_val <= 0:
                    qty_val = float(quantity)

                ws2.row_dimensions[r_idx].height = 19
                ws2.cell(row=r_idx, column=1, value=itm.get("sr_no")).alignment = Alignment(horizontal="center", vertical="center")
                ws2.cell(row=r_idx, column=2, value=itm.get("name")).alignment = Alignment(horizontal="left", vertical="center")

                c3 = ws2.cell(row=r_idx, column=3, value=qty_val)
                c3.alignment = Alignment(horizontal="right", vertical="center")
                c3.number_format = '#,##0.0000' if itm.get("type") != "Shell" else '#,##0'

                c4 = ws2.cell(row=r_idx, column=4, value=rate_val)
                c4.alignment = Alignment(horizontal="right", vertical="center")
                c4.number_format = '#,##0.00'

                c5 = ws2.cell(row=r_idx, column=5, value=f"=C{r_idx}*D{r_idx}")
                c5.alignment = Alignment(horizontal="right", vertical="center")
                c5.number_format = '#,##0.00'

                c6 = ws2.cell(row=r_idx, column=6, value=f"=E{r_idx}/$B$3")
                c6.alignment = Alignment(horizontal="right", vertical="center")
                c6.number_format = '#,##0.0000'

                ws2.cell(row=r_idx, column=7, value=itm.get("group", "")).alignment = Alignment(horizontal="center", vertical="center")

                for c in range(1, 8):
                    ws2.cell(row=r_idx, column=c).border = thin_border
                    ws2.cell(row=r_idx, column=c).font = regular_font
                r_idx += 1

        end_exc_row = r_idx - 1
        tot_exc_row = r_idx

        # Total Excipients Cost Row
        ws2.row_dimensions[tot_exc_row].height = 22
        ws2.merge_cells(start_row=tot_exc_row, start_column=1, end_row=tot_exc_row, end_column=4)
        c_tot_e = ws2.cell(row=tot_exc_row, column=1, value="TOTAL INACTIVE EXCIPIENTS & SHELL COST")
        c_tot_e.font = bold_font
        c_tot_e.alignment = Alignment(horizontal="right", vertical="center")

        if end_exc_row >= start_exc_row:
            c_tot_ev = ws2.cell(row=tot_exc_row, column=5, value=f"=SUM(E{start_exc_row}:E{end_exc_row})")
        else:
            c_tot_ev = ws2.cell(row=tot_exc_row, column=5, value=0.0)
        c_tot_ev.font = bold_font
        c_tot_ev.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_ev.number_format = '#,##0.00'

        c_tot_eu = ws2.cell(row=tot_exc_row, column=6, value=f"=E{tot_exc_row}/$B$3")
        c_tot_eu.font = bold_font
        c_tot_eu.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_eu.number_format = '#,##0.0000'

        for c in range(1, 8):
            cell_obj = ws2.cell(row=tot_exc_row, column=c)
            cell_obj.border = thin_border
            cell_obj.fill = group_fill

        r_idx += 2

        # Section 2: Comprehensive Commercial Cost Derivation Table
        ws2.cell(row=r_idx, column=1, value="2. COMPREHENSIVE COMMERCIAL COST DERIVATION & RATE PER STRIP / TABLET").font = group_font
        r_idx += 1
        deriv_headers = ["Cost Component / Element", "Basis of Calculation", "Total Amount (₹)", f"Cost / Unit ({product_type}) (₹)", pack_col_name]
        ws2.row_dimensions[r_idx].height = 24
        for c_idx, h in enumerate(deriv_headers, 1):
            c = ws2.cell(row=r_idx, column=c_idx, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border

        r_idx += 1
        # Row A: Active Raw Materials Cost
        r_deriv_act = r_idx
        ws2.row_dimensions[r_idx].height = 20
        ws2.cell(row=r_idx, column=1, value="A. Active Raw Materials Cost").font = regular_font
        ws2.cell(row=r_idx, column=2, value="Sum of all active formulation raw materials").font = regular_font
        ws2.cell(row=r_idx, column=3, value=f"=E{tot_act_row}").number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=4, value=f"=C{r_idx}/$B$3").number_format = '#,##0.0000'
        ws2.cell(row=r_idx, column=5, value=f"=D{r_idx}*$D$3").number_format = '#,##0.00'
        for c in range(1, 6):
            ws2.cell(row=r_idx, column=c).border = thin_border
            if c >= 3:
                ws2.cell(row=r_idx, column=c).alignment = Alignment(horizontal="right", vertical="center")
                ws2.cell(row=r_idx, column=c).font = regular_font
        r_idx += 1

        # Row B: Inactive Excipients, Binders & Shells
        r_deriv_exc = r_idx
        ws2.row_dimensions[r_idx].height = 20
        ws2.cell(row=r_idx, column=1, value="B. Inactive Excipients, Binders & Shells").font = regular_font
        ws2.cell(row=r_idx, column=2, value="Granulation, disintegrants, lubricants, coating & shell").font = regular_font
        ws2.cell(row=r_idx, column=3, value=f"=E{tot_exc_row}").number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=4, value=f"=C{r_idx}/$B$3").number_format = '#,##0.0000'
        ws2.cell(row=r_idx, column=5, value=f"=D{r_idx}*$D$3").number_format = '#,##0.00'
        for c in range(1, 6):
            ws2.cell(row=r_idx, column=c).border = thin_border
            if c >= 3:
                ws2.cell(row=r_idx, column=c).alignment = Alignment(horizontal="right", vertical="center")
                ws2.cell(row=r_idx, column=c).font = regular_font
        r_idx += 1

        # Row C: Total Raw Material Cost (A + B)
        r_deriv_rm = r_idx
        ws2.row_dimensions[r_idx].height = 20
        ws2.cell(row=r_idx, column=1, value="C. Total Raw Material Cost (A + B)").font = bold_font
        ws2.cell(row=r_idx, column=2, value="Total raw materials (active + inactive)").font = regular_font
        ws2.cell(row=r_idx, column=3, value=f"=C{r_deriv_act}+C{r_deriv_exc}").number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=4, value=f"=C{r_idx}/$B$3").number_format = '#,##0.0000'
        ws2.cell(row=r_idx, column=5, value=f"=D{r_idx}*$D$3").number_format = '#,##0.00'
        for c in range(1, 6):
            ws2.cell(row=r_idx, column=c).border = thin_border
            if c >= 3:
                ws2.cell(row=r_idx, column=c).alignment = Alignment(horizontal="right", vertical="center")
                ws2.cell(row=r_idx, column=c).font = bold_font
        r_idx += 1

        # Row D: Conversion / Manufacturing Charge
        r_deriv_conv = r_idx
        ws2.row_dimensions[r_idx].height = 20
        ws2.cell(row=r_idx, column=1, value="D. Conversion / Manufacturing Charge").font = regular_font
        ws2.cell(row=r_idx, column=2, value="₹0.25 per unit standard GMP manufacturing charge").font = regular_font
        ws2.cell(row=r_idx, column=4, value=0.25).number_format = '#,##0.0000'
        ws2.cell(row=r_idx, column=3, value=f"=D{r_idx}*$B$3").number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=5, value=f"=D{r_idx}*$D$3").number_format = '#,##0.00'
        for c in range(1, 6):
            ws2.cell(row=r_idx, column=c).border = thin_border
            if c >= 3:
                ws2.cell(row=r_idx, column=c).alignment = Alignment(horizontal="right", vertical="center")
                ws2.cell(row=r_idx, column=c).font = regular_font
        r_idx += 1

        # Row E: Packaging Materials
        r_deriv_pkg = r_idx
        pkg_total_val = float(fin.get("total_packaging_cost", 0.0))
        ws2.row_dimensions[r_idx].height = 20
        ws2.cell(row=r_idx, column=1, value="E. Packaging Materials (Primary + Secondary + Tertiary)").font = regular_font
        ws2.cell(row=r_idx, column=2, value=f"{prim_type} packaging + cartons & master shippers").font = regular_font
        ws2.cell(row=r_idx, column=3, value=pkg_total_val).number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=4, value=f"=C{r_idx}/$B$3").number_format = '#,##0.0000'
        ws2.cell(row=r_idx, column=5, value=f"=D{r_idx}*$D$3").number_format = '#,##0.00'
        for c in range(1, 6):
            ws2.cell(row=r_idx, column=c).border = thin_border
            if c >= 3:
                ws2.cell(row=r_idx, column=c).alignment = Alignment(horizontal="right", vertical="center")
                ws2.cell(row=r_idx, column=c).font = regular_font
        r_idx += 1

        # Row F: Net Manufacturing Cost (C + D + E)
        r_deriv_net = r_idx
        ws2.row_dimensions[r_idx].height = 20
        ws2.cell(row=r_idx, column=1, value="F. Net Manufacturing Cost (C + D + E)").font = bold_font
        ws2.cell(row=r_idx, column=2, value="Total prime manufacturing cost before operating margin").font = regular_font
        ws2.cell(row=r_idx, column=3, value=f"=C{r_deriv_rm}+C{r_deriv_conv}+C{r_deriv_pkg}").number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=4, value=f"=C{r_idx}/$B$3").number_format = '#,##0.0000'
        ws2.cell(row=r_idx, column=5, value=f"=D{r_idx}*$D$3").number_format = '#,##0.00'
        for c in range(1, 6):
            ws2.cell(row=r_idx, column=c).border = thin_border
            if c >= 3:
                ws2.cell(row=r_idx, column=c).alignment = Alignment(horizontal="right", vertical="center")
                ws2.cell(row=r_idx, column=c).font = bold_font
        r_idx += 1

        # Row G: Walpar Operating Margin (20%)
        r_deriv_margin = r_idx
        ws2.row_dimensions[r_idx].height = 20
        ws2.cell(row=r_idx, column=1, value="G. Walpar Operating Margin (20%)").font = regular_font
        ws2.cell(row=r_idx, column=2, value="20% operating margin on net manufacturing cost").font = regular_font
        ws2.cell(row=r_idx, column=3, value=f"=C{r_deriv_net}*0.20").number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=4, value=f"=C{r_idx}/$B$3").number_format = '#,##0.0000'
        ws2.cell(row=r_idx, column=5, value=f"=D{r_idx}*$D$3").number_format = '#,##0.00'
        for c in range(1, 6):
            ws2.cell(row=r_idx, column=c).border = thin_border
            if c >= 3:
                ws2.cell(row=r_idx, column=c).alignment = Alignment(horizontal="right", vertical="center")
                ws2.cell(row=r_idx, column=c).font = regular_font
        r_idx += 1

        # Row H: TOTAL COMMERCIAL BATCH VALUE (F + G)
        r_deriv_final = r_idx
        ws2.row_dimensions[r_idx].height = 24
        ws2.cell(row=r_idx, column=1, value="H. TOTAL COMMERCIAL BATCH VALUE (F + G)").font = bold_font
        ws2.cell(row=r_idx, column=2, value=f"Final commercial quotation for {quantity:,} units").font = regular_font
        ws2.cell(row=r_idx, column=3, value=f"=C{r_deriv_net}+C{r_deriv_margin}").number_format = '#,##0.00'
        ws2.cell(row=r_idx, column=4, value=f"=C{r_idx}/$B$3").number_format = '#,##0.0000'
        ws2.cell(row=r_idx, column=5, value=f"=D{r_idx}*$D$3").number_format = '#,##0.00'
        for c in range(1, 6):
            cell_obj = ws2.cell(row=r_idx, column=c)
            cell_obj.border = thin_border
            cell_obj.fill = rate_fill
            if c >= 3:
                cell_obj.alignment = Alignment(horizontal="right", vertical="center")
                cell_obj.font = bold_font
        r_idx += 2

        # Section 3: Official Commercial Pricing KPI Cards (Live Formulas)
        ws2.row_dimensions[r_idx].height = 24
        ws2.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=5)
        hdr_kpi = ws2.cell(row=r_idx, column=1, value="3. OFFICIAL COMMERCIAL PRICING SUMMARY (DYNAMIC REAL-TIME FORMULAS)")
        hdr_kpi.font = group_font
        hdr_kpi.fill = gold_fill
        hdr_kpi.alignment = Alignment(horizontal="center", vertical="center")
        for c in range(1, 6):
            ws2.cell(row=r_idx, column=c).border = thin_border
        r_idx += 1

        # KPI 1: Rate Per Unit (Tablet / Capsule / Bottle)
        r_kpi_unit = r_idx
        ws2.row_dimensions[r_idx].height = 26
        ws2.merge_cells(start_row=r_kpi_unit, start_column=1, end_row=r_kpi_unit, end_column=3)
        c1 = ws2.cell(row=r_kpi_unit, column=1, value=f"* FINAL RATE PER {product_type.upper()} (PER UNIT):")
        c1.font = kpi_title_font
        c1.alignment = Alignment(horizontal="left", vertical="center")

        ws2.merge_cells(start_row=r_kpi_unit, start_column=4, end_row=r_kpi_unit, end_column=5)
        v1 = ws2.cell(row=r_kpi_unit, column=4, value=f"=D{r_deriv_final}")
        v1.font = kpi_val_font
        v1.alignment = Alignment(horizontal="center", vertical="center")
        v1.number_format = '"₹"#,##0.0000'
        v1.fill = rate_fill
        for c in range(1, 6):
            ws2.cell(row=r_kpi_unit, column=c).border = thin_border
        r_idx += 1

        # KPI 2: Rate Per Strip / Jar / Pack
        r_kpi_pack = r_idx
        ws2.row_dimensions[r_idx].height = 28
        ws2.merge_cells(start_row=r_kpi_pack, start_column=1, end_row=r_kpi_pack, end_column=3)
        c2 = ws2.cell(row=r_kpi_pack, column=1, value=f"* FINAL RATE PER {pack_label} ({pack_size_val} {product_type.upper()}S):")
        c2.font = kpi_title_font
        c2.alignment = Alignment(horizontal="left", vertical="center")

        ws2.merge_cells(start_row=r_kpi_pack, start_column=4, end_row=r_kpi_pack, end_column=5)
        v2 = ws2.cell(row=r_kpi_pack, column=4, value=f"=E{r_deriv_final}")
        v2.font = Font(name="Calibri", size=14, bold=True, color="0F5132")
        v2.alignment = Alignment(horizontal="center", vertical="center")
        v2.number_format = '"₹"#,##0.00'
        v2.fill = rate_fill
        for c in range(1, 6):
            ws2.cell(row=r_kpi_pack, column=c).border = thin_border
        r_idx += 1

        # KPI 3: Total Commercial Batch Value
        r_kpi_tot = r_idx
        ws2.row_dimensions[r_idx].height = 26
        ws2.merge_cells(start_row=r_kpi_tot, start_column=1, end_row=r_kpi_tot, end_column=3)
        c3 = ws2.cell(row=r_kpi_tot, column=1, value=f"* TOTAL COMMERCIAL BATCH VALUE ({quantity:,} UNITS):")
        c3.font = kpi_title_font
        c3.alignment = Alignment(horizontal="left", vertical="center")

        ws2.merge_cells(start_row=r_kpi_tot, start_column=4, end_row=r_kpi_tot, end_column=5)
        v3 = ws2.cell(row=r_kpi_tot, column=4, value=f"=C{r_deriv_final}")
        v3.font = kpi_tot_font
        v3.alignment = Alignment(horizontal="center", vertical="center")
        v3.number_format = '"₹"#,##0.00'
        v3.fill = gold_fill
        for c in range(1, 6):
            ws2.cell(row=r_kpi_tot, column=c).border = thin_border
        r_idx += 1

        # Formula Notice Callout
        ws2.row_dimensions[r_idx].height = 24
        ws2.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=7)
        note = ws2.cell(row=r_idx, column=1, value="💡 DYNAMIC SPREADSHEET NOTICE: All rates and values in this workbook are linked dynamically via Excel formulas. If you modify any ingredient rate in Column D above, the Total Cost, Manufacturing Cost, Rate Per Unit, and Rate Per Strip will recalculate automatically in real time.")
        note.font = Font(name="Calibri", size=9, italic=True, color="444444")
        note.alignment = Alignment(horizontal="left", vertical="center")

        # Cross reference commercial rates into Sheet 1 technical summary
        ws.merge_cells(start_row=current_row, start_column=5, end_row=current_row, end_column=7)
        lbl_ru = ws.cell(row=current_row, column=5, value=f"Commercial Rate / {product_type}:")
        lbl_ru.font = bold_font
        lbl_ru.alignment = Alignment(horizontal="right", vertical="center")
        lbl_ru.fill = rate_fill

        ws.merge_cells(start_row=current_row, start_column=8, end_row=current_row, end_column=9)
        val_ru = ws.cell(row=current_row, column=8, value=f"='Costing & Commercial Rates'!D{r_deriv_final}")
        val_ru.font = bold_font
        val_ru.alignment = Alignment(horizontal="left", vertical="center")
        val_ru.fill = rate_fill
        val_ru.number_format = '"₹"#,##0.0000'
        for c in range(5, 10):
            ws.cell(row=current_row, column=c).border = thin_border
        current_row += 1

        ws.merge_cells(start_row=current_row, start_column=5, end_row=current_row, end_column=7)
        lbl_rs = ws.cell(row=current_row, column=5, value=f"Commercial Rate / {pack_label}:")
        lbl_rs.font = bold_font
        lbl_rs.alignment = Alignment(horizontal="right", vertical="center")
        lbl_rs.fill = rate_fill

        ws.merge_cells(start_row=current_row, start_column=8, end_row=current_row, end_column=9)
        val_rs = ws.cell(row=current_row, column=8, value=f"='Costing & Commercial Rates'!E{r_deriv_final}")
        val_rs.font = bold_font
        val_rs.alignment = Alignment(horizontal="left", vertical="center")
        val_rs.fill = rate_fill
        val_rs.number_format = '"₹"#,##0.00'
        for c in range(5, 10):
            ws.cell(row=current_row, column=c).border = thin_border

        # Adjust column widths for Sheet 2
        for col in ws2.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and not cell.coordinate in ws2.merged_cells:
                    max_len = len(val_str)
            ws2.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def get_all_ingredient_rates(self) -> List[Dict[str, Any]]:
        """Return all ingredients with current rates sorted alphabetically"""
        self._load_rate_data()
        results = []
        for name in sorted(self.rate_dict.keys()):
            results.append({
                "name": name,
                "rate": round(self.rate_dict[name], 2)
            })
        return results

    def update_rates_from_file(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        """
        Upload new ingredient rate sheet (.xlsx, .xls, .csv).
        Detect columns, update existing rates, find similar names using intelligent matching,
        and add new ingredients. Saves permanently to rate avg.xlsx with backup.
        """
        import difflib
        import io
        import re
        import shutil
        from datetime import datetime
        from pathlib import Path

        fn_lower = filename.lower()
        try:
            if fn_lower.endswith(".csv"):
                df = pd.read_csv(io.BytesIO(file_content))
            else:
                df = pd.read_excel(io.BytesIO(file_content))
        except Exception as read_err:
            return {"success": False, "error": f"Failed to parse file: {str(read_err)}"}

        if df.empty:
            return {"success": False, "error": "Uploaded rate sheet is empty."}

        # Identify ingredient and rate columns
        item_col = None
        rate_col = None

        for col in df.columns:
            c_str = str(col).lower().strip()
            if any(k in c_str for k in ["ingredient", "item", "raw material", "material", "name", "particular"]):
                if item_col is None:
                    item_col = col
            elif any(k in c_str for k in ["rate", "price", "cost", "rs", "₹", "amount"]):
                if rate_col is None:
                    rate_col = col

        if item_col is None and len(df.columns) >= 1:
            item_col = df.columns[0]
        if rate_col is None and len(df.columns) >= 2:
            rate_col = df.columns[1]

        if item_col is None or rate_col is None:
            return {"success": False, "error": "Could not identify Ingredient Name and Rate columns in uploaded sheet."}

        # Backup existing rate file
        backup_dir = Path("app/data/rate_backups")
        backup_dir.mkdir(parents=True, exist_ok=True)
        if os.path.exists(self.rate_file_path):
            timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_path = backup_dir / f"rate_avg_backup_{timestamp_str}.xlsx"
            try:
                shutil.copyfile(self.rate_file_path, backup_path)
            except Exception as bkp_err:
                print(f"[Rate Update] Backup notice: {bkp_err}")

        # Ensure current rate_dict is up to date
        self._load_rate_data()

        total_processed = 0
        exact_updated = 0
        similar_updated = 0
        new_added = 0
        sample_updates = []

        existing_keys = list(self.rate_dict.keys())

        def clean_name(val: Any) -> str:
            s = str(val).strip().lower()
            s = re.sub(r'[\r\n\t]+', ' ', s)
            s = re.sub(r'\s+', ' ', s)
            s = s.strip(' "\'.,:;()[]')
            return s

        for _, row in df.iterrows():
            raw_item = row[item_col]
            raw_rate = row[rate_col]

            if pd.isna(raw_item) or pd.isna(raw_rate):
                continue

            c_name = clean_name(raw_item)
            if not c_name or c_name in ["item", "ingredient", "material", "name", "total", "sr", "no"]:
                continue

            try:
                clean_rate_str = str(raw_rate).replace(",", "").replace("₹", "").replace("/-", "").strip()
                rate_val = float(clean_rate_str)
            except Exception:
                continue

            if rate_val <= 0:
                continue

            total_processed += 1

            # 1. Exact match check
            if c_name in self.rate_dict:
                old_rate = self.rate_dict[c_name]
                self.rate_dict[c_name] = rate_val
                exact_updated += 1
                if len(sample_updates) < 20:
                    sample_updates.append({
                        "item": c_name,
                        "matched_to": c_name,
                        "match_type": "Exact Match",
                        "old_rate": round(old_rate, 2),
                        "new_rate": round(rate_val, 2)
                    })
                continue

            # 2. Smart Fuzzy Match
            close_matches = difflib.get_close_matches(c_name, existing_keys, n=1, cutoff=0.75)
            if close_matches:
                matched_key = close_matches[0]
                old_rate = self.rate_dict[matched_key]
                self.rate_dict[matched_key] = rate_val
                similar_updated += 1
                if len(sample_updates) < 20:
                    sample_updates.append({
                        "item": c_name,
                        "matched_to": matched_key,
                        "match_type": "Intelligent Similarity Match",
                        "old_rate": round(old_rate, 2),
                        "new_rate": round(rate_val, 2)
                    })
                continue

            # 3. New ingredient
            self.rate_dict[c_name] = rate_val
            existing_keys.append(c_name)
            new_added += 1
            if len(sample_updates) < 20:
                sample_updates.append({
                    "item": c_name,
                    "matched_to": "(New Entry)",
                    "match_type": "New Ingredient Added",
                    "old_rate": 0.0,
                    "new_rate": round(rate_val, 2)
                })

        # Save to rate avg.xlsx
        output_rows = [{"Item": k, "Rate": v} for k, v in sorted(self.rate_dict.items())]
        save_df = pd.DataFrame(output_rows)
        save_df.to_excel(self.rate_file_path, index=False)

        # Refresh in-memory list
        self._load_rate_data()

        return {
            "success": True,
            "total_processed": total_processed,
            "exact_updated": exact_updated,
            "similar_updated": similar_updated,
            "new_added": new_added,
            "sample_updates": sample_updates,
            "total_database_items": len(self.rate_dict)
        }

    def update_single_rate(self, ingredient_name: str, new_rate: float) -> Dict[str, Any]:
        """
        Directly update or add a single ingredient rate from the Admin panel,
        create backup of rate avg.xlsx, save to Excel, and reload cache.
        """
        if not ingredient_name or not str(ingredient_name).strip():
            return {"success": False, "error": "Ingredient name cannot be empty"}
        try:
            rate_val = float(new_rate)
            if rate_val < 0:
                return {"success": False, "error": "Rate cannot be negative"}
        except Exception:
            return {"success": False, "error": "Invalid rate value"}

        # Backup rate avg.xlsx
        rate_path = Path(self.rate_file_path)
        if rate_path.exists():
            backup_dir = rate_path.parent / "rate_backups"
            backup_dir.mkdir(parents=True, exist_ok=True)
            timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_path = backup_dir / f"rate_avg_backup_{timestamp_str}.xlsx"
            try:
                shutil.copyfile(str(rate_path), str(backup_path))
            except Exception as bkp_err:
                print(f"[Rate Update] Backup notice: {bkp_err}")

        # Ensure rate_dict is up to date
        self._load_rate_data()

        clean_name = str(ingredient_name).strip().lower()
        clean_name = re.sub(r'[\r\n\t]+', ' ', clean_name)
        clean_name = re.sub(r'\s+', ' ', clean_name)
        clean_name = clean_name.strip(' "\'.,:;()[]')

        old_rate = 0.0
        is_update = False
        target_key = clean_name

        if clean_name in self.rate_dict:
            target_key = clean_name
            old_rate = self.rate_dict[clean_name]
            is_update = True
        else:
            close_matches = difflib.get_close_matches(clean_name, list(self.rate_dict.keys()), n=1, cutoff=0.82)
            if close_matches:
                target_key = close_matches[0]
                old_rate = self.rate_dict[target_key]
                is_update = True

        self.rate_dict[target_key] = rate_val

        # Save to rate avg.xlsx
        output_rows = [{"Item": k, "Rate": v} for k, v in sorted(self.rate_dict.items())]
        save_df = pd.DataFrame(output_rows)
        save_df.to_excel(self.rate_file_path, index=False)

        # Reload cache
        self._load_rate_data()

        return {
            "success": True,
            "ingredient": target_key,
            "original_query": ingredient_name,
            "is_update": is_update,
            "old_rate": round(old_rate, 2),
            "new_rate": round(rate_val, 2),
            "total_count": len(self.rate_dict)
        }
