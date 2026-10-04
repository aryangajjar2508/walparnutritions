import os
import re
import sys
import shutil
import uuid
import cv2
import numpy as np
from pathlib import Path
from typing import List, Dict, Any, Optional
from PIL import Image, ImageOps

# ── Force UTF-8 on Windows (prevents charmap errors from OCR special chars) ──
os.environ.setdefault('PYTHONIOENCODING', 'utf-8')
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass
if hasattr(sys.stderr, 'reconfigure'):
    try:
        sys.stderr.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

try:
    import pillow_heif
    pillow_heif.register_heif_opener()
except Exception:
    pass

try:
    import fitz
except Exception:
    pass

from rapidfuzz import fuzz

from fastapi import FastAPI, File, UploadFile, Request, Form, Query
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel

from app.config import (
    UPLOAD_DIR, SAMPLE_DIR, STATIC_DIR, TEMPLATES_DIR, CLOUDFLARE_URL_FILE
)
from app.ocr_manager import OCRManager
from app.formula_parser import FormulaParser
from app.learning_engine import SelfLearningEngine
from app.evaluator import FormulaEvaluator
from app.batch_master import BatchMasterEngine
from app.salt_rda_engine import SaltRDAEngine
from app.auth import (
    authenticate_user, create_session_token, verify_session_token,
    log_login_event, log_upload_event, log_batch_download_event,
    get_all_users, create_user, update_user_password, toggle_user_active,
    delete_user, get_upload_logs, get_login_logs, get_batch_logs,
    get_current_ist_str, MASTER_BMR_PASSWORD,
    log_quotation_event, get_pending_missing_rates, get_all_missing_rate_requests,
    resolve_missing_rate_requests, get_quotation_by_id, update_quotation_revised_rates,
    get_per_user_quotation_stats, log_user_estimated_rate, get_user_estimated_rates
)
from app.master_config import master_config_mgr

app = FastAPI(title="Walpar Neural Formula OCR & Formulation AI", version="3.2.0")

# Mount static, uploads, and samples directories
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
app.mount("/uploads", StaticFiles(directory=str(UPLOAD_DIR)), name="uploads")
app.mount("/samples", StaticFiles(directory=str(SAMPLE_DIR)), name="samples")

templates = Jinja2Templates(directory=str(TEMPLATES_DIR))

def render_template(request: Request, name: str, context: Optional[Dict[str, Any]] = None, status_code: int = 200) -> HTMLResponse:
    """Universal template response helper supporting both Starlette 0.36+ and legacy versions"""
    ctx = dict(context or {})
    ctx["request"] = request
    try:
        # Starlette 0.36+ keyword syntax
        return templates.TemplateResponse(request=request, name=name, context=ctx, status_code=status_code)
    except TypeError:
        # Legacy Starlette positional syntax
        return templates.TemplateResponse(name, ctx, status_code=status_code)

# Initialize services
ocr_manager = OCRManager.get_instance()
formula_parser = FormulaParser()
learning_engine = SelfLearningEngine.get_instance()
evaluator = FormulaEvaluator()
batch_master_engine = BatchMasterEngine.get_instance()
salt_rda_engine = SaltRDAEngine.get_instance()

def get_current_user(request: Request) -> Optional[Dict[str, Any]]:
    cookie = request.cookies.get("walpar_session")
    return verify_session_token(cookie)

class IngredientItem(BaseModel):
    id: Optional[str] = None
    name: str
    dosage: float
    unit: str = "mg"
    db_raw_name: Optional[str] = ""
    is_matched: Optional[bool] = False
    match_score: Optional[float] = 0.0

class EvaluatePayload(BaseModel):
    ingredients: List[Dict[str, Any]]
    dosage_form: Optional[str] = "tablets"

# ══════════════════════════════════════════════════════════════════
# AUTHENTICATION & ACCESS CONTROL (ID / PASSWORD WALL)
# ══════════════════════════════════════════════════════════════════

@app.get("/api/health")
async def health_check():
    return {"status": "ok", "app": "Walpar Neural Formula OCR", "version": "3.2.0"}

@app.get("/login", response_class=HTMLResponse)
async def login_page(request: Request):
    user = get_current_user(request)
    if user:
        return RedirectResponse(url="/", status_code=303)
    error = request.query_params.get("error")
    return render_template(request, "login.html", {
        "error": error
    })

@app.post("/api/auth/login")
async def api_auth_login(
    request: Request,
    username: str = Form(...),
    password: str = Form(...)
):
    ip = request.client.host if request.client else "127.0.0.1"
    forwarded = request.headers.get("x-forwarded-for")
    if forwarded:
        ip = forwarded.split(",")[0].strip()
    user_agent = request.headers.get("user-agent", "Unknown Device")

    user = authenticate_user(username, password)
    if not user:
        log_login_event(username, ip, user_agent, status="FAILED")
        return RedirectResponse(url="/login?error=Invalid+User+ID+or+Password", status_code=303)

    log_login_event(user["username"], ip, user_agent, status="SUCCESS")
    token = create_session_token(user["username"], user["role"])

    response = RedirectResponse(url="/", status_code=303)
    response.set_cookie(
        key="walpar_session",
        value=token,
        max_age=14 * 86400,
        httponly=True,
        samesite="lax",
        path="/"
    )
    return response

@app.get("/api/auth/logout")
async def api_auth_logout():
    response = RedirectResponse(url="/login", status_code=303)
    response.delete_cookie(key="walpar_session", path="/")
    return response

@app.get("/admin", response_class=HTMLResponse)
async def admin_panel(request: Request):
    user = get_current_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=303)
    if user.get("role") != "admin":
        return RedirectResponse(url="/?error=Administrator+privileges+required", status_code=303)

    users = get_all_users()
    uploads = get_upload_logs(limit=200)
    logins = get_login_logs(limit=200)
    batch_logs = get_batch_logs(limit=200)
    all_rates = batch_master_engine.get_all_ingredient_rates()
    pending_missing_rates = get_pending_missing_rates()
    all_missing_rates = get_all_missing_rate_requests(limit=100)
    user_estimated_rates = get_user_estimated_rates(limit=100)
    per_user_stats = get_per_user_quotation_stats()
    total_pipeline_val = sum(float(x.get("total_batch_val") or 0.0) for x in batch_logs)

    return render_template(request, "admin.html", {
        "current_user": user,
        "users": users,
        "uploads": uploads,
        "logins": logins,
        "batch_logs": batch_logs,
        "total_pipeline_val": total_pipeline_val,
        "all_rates": all_rates,
        "pending_missing_rates": pending_missing_rates,
        "all_missing_rates": all_missing_rates,
        "user_estimated_rates": user_estimated_rates,
        "per_user_stats": per_user_stats,
        "master_config": master_config_mgr.get_config()
    })

@app.post("/api/admin/users/create")
async def api_admin_create_user(request: Request):
    user = get_current_user(request)
    if not user or user.get("role") != "admin":
        return JSONResponse({"success": False, "error": "Unauthorized"}, status_code=403)
    body = await request.json()
    res = create_user(
        username=body.get("username", ""),
        password=body.get("password", ""),
        full_name=body.get("full_name", ""),
        role=body.get("role", "user")
    )
    return JSONResponse(res)

@app.post("/api/admin/users/reset-password")
async def api_admin_reset_password(request: Request):
    user = get_current_user(request)
    if not user or user.get("role") != "admin":
        return JSONResponse({"success": False, "error": "Unauthorized"}, status_code=403)
    body = await request.json()
    res = update_user_password(body.get("username", ""), body.get("new_password", ""))
    return JSONResponse(res)

@app.post("/api/admin/users/toggle-status")
async def api_admin_toggle_status(request: Request):
    user = get_current_user(request)
    if not user or user.get("role") != "admin":
        return JSONResponse({"success": False, "error": "Unauthorized"}, status_code=403)
    body = await request.json()
    res = toggle_user_active(body.get("username", ""))
    return JSONResponse(res)

@app.post("/api/admin/users/delete")
async def api_admin_delete_user(request: Request):
    user = get_current_user(request)
    if not user or user.get("role") != "admin":
        return JSONResponse({"success": False, "error": "Unauthorized"}, status_code=403)
    body = await request.json()
    res = delete_user(body.get("username", ""))
    return JSONResponse(res)

# ══════════════════════════════════════════════════════════════════
# MAIN OCR DASHBOARD (REQUIRES LOGIN)
# ══════════════════════════════════════════════════════════════════

@app.get("/", response_class=HTMLResponse)
async def index_page(request: Request):
    user = get_current_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=303)

    samples = [
        {
            "id": "sample_immunity_booster.png",
            "name": "Immunity Booster & Defense",
            "subtitle": "Vitamin C, Zinc, D3, Elderberry"
        },
        {
            "id": "sample_joint_care.png",
            "name": "Advanced Joint Mobility Matrix",
            "subtitle": "Glucosamine, Chondroitin, MSM, Boswellia"
        },
        {
            "id": "sample_curcumin_synergy.png",
            "name": "Bio-Enhanced Curcumin",
            "subtitle": "Curcumin 95% + Piperine 95%"
        },
        {
            "id": "sample_sleep_relax.png",
            "name": "Deep Sleep & Relaxation",
            "subtitle": "Melatonin, L-Theanine, Magnesium, B6"
        }
    ]

    import socket
    local_ip = "127.0.0.1"
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(('8.8.8.8', 80))
        local_ip = s.getsockname()[0]
        s.close()
    except Exception:
        pass
    port = request.url.port or 8000

    cloudflare_url = os.environ.get("CLOUDFLARE_PUBLIC_URL")
    if not cloudflare_url and CLOUDFLARE_URL_FILE.exists():
        try:
            val = CLOUDFLARE_URL_FILE.read_text(encoding="utf-8").strip()
            if val.startswith("https://") and ".trycloudflare.com" in val:
                cloudflare_url = val
        except Exception:
            pass

    mobile_url = cloudflare_url if cloudflare_url else f"http://{local_ip}:{port}"
    is_cloudflare = bool(cloudflare_url)

    return render_template(request, "index.html", {
        "current_user": user,
        "samples": samples,
        "total_db_ingredients": len(formula_parser.ingredients_master),
        "ocr_engines": ocr_manager.get_available_engines(),
        "learning_status": learning_engine.get_learning_status(),
        "mobile_url": mobile_url,
        "local_ip": local_ip,
        "port": port,
        "is_cloudflare": is_cloudflare,
        "cloudflare_url": cloudflare_url
    })

@app.get("/api/engines")
async def get_engines():
    return {"engines": ocr_manager.get_available_engines()}

@app.get("/api/learning-status")
async def get_learning_status():
    return learning_engine.get_learning_status()

@app.get("/api/ingredients")
async def get_master_ingredients(q: Optional[str] = None):
    items = []
    for item in formula_parser.ingredients_master:
        rate = float(item.get("_rate") or 0.0)
        r_info = {}
        if rate <= 0:
            r_info = batch_master_engine.get_rate_info(item["name"])
            if not r_info.get("available") and item.get("raw_name"):
                r_info = batch_master_engine.get_rate_info(item["raw_name"])
            rate = float(r_info.get("rate", 0.0))

        item_data = {
            "id": item["id"],
            "name": item["name"],
            "raw_name": item.get("raw_name", ""),
            "default_unit": item.get("default_unit", "mg"),
            "default_dose": item.get("default_dose", 100),
            "aliases": item.get("aliases", []),
            "rate": rate,
            "is_rate_available": rate > 0,
            "rate_matched_name": r_info.get("matched_name", item.get("raw_name", item["name"]))
        }
        if q:
            query = q.lower().strip()
            if query in item["name"].lower() or any(query in a.lower() for a in item.get("aliases", [])):
                items.append(item_data)
        else:
            items.append(item_data)

    return {"ingredients": items, "total": len(items)}

@app.post("/api/evaluate")
async def evaluate_formula(payload: EvaluatePayload):
    try:
        evaluation = evaluator.evaluate_formula(payload.ingredients, payload.dosage_form or "tablets")
        return JSONResponse({"success": True, "evaluation": evaluation})
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.post("/api/detect-corners")
async def api_detect_corners(
    file: Optional[UploadFile] = File(None),
    image_url: Optional[str] = Form(None)
):
    try:
        from app.image_preprocessor import detect_document_corners
        target_path = None
        temp_created = False

        if file and file.filename:
            temp_name = f"temp_detect_{uuid.uuid4().hex[:8]}.jpg"
            target_path = UPLOAD_DIR / temp_name
            with open(target_path, "wb") as buf:
                shutil.copyfileobj(file.file, buf)
            temp_created = True
        elif image_url:
            clean_url = image_url.split("?")[0].lstrip("/")
            if clean_url.startswith("uploads/"):
                target_path = UPLOAD_DIR / clean_url.replace("uploads/", "")
            elif clean_url.startswith("samples/"):
                target_path = SAMPLE_DIR / clean_url.replace("samples/", "")

        if not target_path or not Path(target_path).exists():
            return JSONResponse({
                "success": False,
                "error": "Image not found for corner detection"
            }, status_code=400)

        corners = detect_document_corners(target_path)

        if temp_created and target_path and Path(target_path).exists():
            try:
                os.remove(target_path)
            except Exception:
                pass

        return JSONResponse({"success": True, "corners": corners})
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.post("/api/crop-perspective")
async def api_crop_perspective(
    file: Optional[UploadFile] = File(None),
    corners: str = Form(...),
    rotation: int = Form(0),
    image_url: Optional[str] = Form(None)
):
    try:
        import json
        from app.image_preprocessor import warp_perspective_quad
        clean_c = corners.strip()
        if "'" in clean_c and '"' not in clean_c:
            clean_c = clean_c.replace("'", '"')
        parsed_corners = json.loads(clean_c)

        target_path = None
        temp_created = False

        if file and file.filename:
            temp_name = f"temp_crop_{uuid.uuid4().hex[:8]}.jpg"
            target_path = UPLOAD_DIR / temp_name
            with open(target_path, "wb") as buf:
                shutil.copyfileobj(file.file, buf)
            temp_created = True
        elif image_url:
            clean_url = image_url.split("?")[0].lstrip("/")
            if clean_url.startswith("uploads/"):
                target_path = UPLOAD_DIR / clean_url.replace("uploads/", "")
            elif clean_url.startswith("samples/"):
                target_path = SAMPLE_DIR / clean_url.replace("samples/", "")

        if not target_path or not Path(target_path).exists():
            return JSONResponse({
                "success": False,
                "error": "Source image not found for perspective unwarp"
            }, status_code=400)

        warped_filename = f"warped_{uuid.uuid4().hex[:10]}.jpg"
        warped_dest = UPLOAD_DIR / warped_filename

        warp_perspective_quad(target_path, parsed_corners, rotation=rotation, save_path=warped_dest)

        if temp_created and target_path and Path(target_path).exists():
            try:
                os.remove(target_path)
            except Exception:
                pass

        return JSONResponse({
            "success": True,
            "image_url": f"/uploads/{warped_filename}",
            "filename": warped_filename
        })
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.post("/api/upload")
async def upload_and_process(
    request: Request,
    file: UploadFile = File(...),
    engine: str = Form("gemini"),
    crop_corners: Optional[str] = Form(None),
    rotation: int = Form(0)
):
    """
    Upload customer formula photo, run doc-scanner perspective unwarp if crop_corners
    is provided, run selected OCR engine, match ingredients, annotate salt & elemental RDA,
    and audit log user and software rates in IST.
    """
    try:
        user = get_current_user(request)
        username = user["username"] if user else "guest"

        orig_name = file.filename or "upload.jpg"
        clean_ext = Path(orig_name).suffix.lower()
        valid_extensions = [
            ".jpg", ".jpeg", ".png", ".webp", ".bmp", ".jfif",
            ".tiff", ".tif", ".pdf", ".heic", ".heif",
            ".docx", ".doc", ".xlsx", ".xls", ".csv", ".txt"
        ]
        if not clean_ext or clean_ext not in valid_extensions:
            clean_ext = ".jpg"
        
        safe_stem = re.sub(r'[^a-zA-Z0-9_-]', '_', Path(orig_name).stem)[:30]
        unique_base = f"{uuid.uuid4().hex[:12]}_{safe_stem}"
        unique_filename = f"{unique_base}{clean_ext}"
        dest_path = UPLOAD_DIR / unique_filename

        with open(dest_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)

        document_extracted_lines = []
        is_document_upload = clean_ext in [".pdf", ".docx", ".doc", ".xlsx", ".xls", ".csv", ".txt"]

        # ── Step 0: Convert Document (PDF, Word, Excel, CSV, TXT) or HEIC to standard JPEG ──
        if is_document_upload:
            doc_rendered_filename = f"{unique_base}.jpg"
            rendered_dest = UPLOAD_DIR / doc_rendered_filename
            try:
                if clean_ext == ".pdf":
                    import fitz
                    doc = fitz.open(str(dest_path))
                    num_pages = len(doc)
                    page_images = []
                    for p_idx in range(min(num_pages, 6)):
                        page = doc[p_idx]
                        t_c = page.get_text("text")
                        if t_c and t_c.strip():
                            for raw_l in t_c.splitlines():
                                clean_l = raw_l.strip()
                                if clean_l and len(clean_l) > 1:
                                    document_extracted_lines.append(clean_l)
                        pix = page.get_pixmap(dpi=300)
                        p_img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                        page_images.append(p_img)
                    doc.close()

                    if page_images:
                        if len(page_images) == 1:
                            page_images[0].save(rendered_dest, "JPEG", quality=95)
                        else:
                            total_w = max(img.width for img in page_images)
                            total_h = sum(img.height for img in page_images) + (len(page_images) - 1) * 20
                            stitched = Image.new("RGB", (total_w, total_h), (255, 255, 255))
                            cur_y = 0
                            for img in page_images:
                                stitched.paste(img, (0, cur_y))
                                cur_y += img.height + 20
                            stitched.save(rendered_dest, "JPEG", quality=95)
                    else:
                        Image.new("RGB", (800, 1000), (255, 255, 255)).save(rendered_dest, "JPEG")

                elif clean_ext in [".docx", ".doc"]:
                    try:
                        import docx
                        d = docx.Document(str(dest_path))
                        for p in d.paragraphs:
                            txt = p.text.strip()
                            if txt:
                                document_extracted_lines.append(txt)
                        for t in d.tables:
                            for r in t.rows:
                                row_t = " ".join(c.text.strip() for c in r.cells if c.text.strip())
                                if row_t:
                                    document_extracted_lines.append(row_t)
                    except Exception as docx_err:
                        print(f"[Upload] Word read notice: {docx_err}")
                    
                    from app.image_preprocessor import render_text_document_image
                    doc_img = render_text_document_image(document_extracted_lines, title=f"WORD FORMULATION • {orig_name}")
                    doc_img.save(rendered_dest, "JPEG", quality=95)

                elif clean_ext in [".xlsx", ".xls", ".csv"]:
                    try:
                        if clean_ext == ".csv":
                            import csv
                            with open(str(dest_path), "r", encoding="utf-8", errors="ignore") as f:
                                for r in csv.reader(f):
                                    row_txt = " ".join(c.strip() for c in r if c.strip())
                                    if row_txt:
                                        document_extracted_lines.append(row_txt)
                        else:
                            import openpyxl
                            wb = openpyxl.load_workbook(str(dest_path), data_only=True)
                            for s in wb.sheetnames:
                                ws = wb[s]
                                for r in ws.iter_rows(values_only=True):
                                    row_cells = [str(x).strip() for x in r if x is not None and str(x).strip()]
                                    if row_cells:
                                        document_extracted_lines.append(" ".join(row_cells))
                    except Exception as xl_err:
                        print(f"[Upload] Excel read notice: {xl_err}")

                    from app.image_preprocessor import render_text_document_image
                    doc_img = render_text_document_image(document_extracted_lines, title=f"EXCEL FORMULATION • {orig_name}")
                    doc_img.save(rendered_dest, "JPEG", quality=95)

                elif clean_ext == ".txt":
                    try:
                        with open(str(dest_path), "r", encoding="utf-8", errors="ignore") as f:
                            document_extracted_lines = [l.strip() for l in f if l.strip()]
                    except Exception as txt_err:
                        print(f"[Upload] Text read notice: {txt_err}")
                    
                    from app.image_preprocessor import render_text_document_image
                    doc_img = render_text_document_image(document_extracted_lines, title=f"TEXT FORMULATION • {orig_name}")
                    doc_img.save(rendered_dest, "JPEG", quality=95)

                dest_path = rendered_dest
                unique_filename = doc_rendered_filename
            except Exception as conv_err:
                print(f"[Upload] Document conversion error: {conv_err}")

        elif clean_ext in [".heic", ".heif"]:
            try:
                heic_rendered_filename = f"{unique_base}.jpg"
                rendered_dest = UPLOAD_DIR / heic_rendered_filename
                with Image.open(dest_path) as im:
                    im.convert("RGB").save(rendered_dest, "JPEG", quality=95)
                dest_path = rendered_dest
                unique_filename = heic_rendered_filename
            except Exception as heic_err:
                print(f"[Upload] HEIC conversion error: {heic_err}")

        # ── Step 1: Doc Scanner Perspective Unwarp ──
        crop_applied = False
        if crop_corners:
            try:
                clean_c = crop_corners.strip()
                if "'" in clean_c and '"' not in clean_c:
                    clean_c = clean_c.replace("'", '"')
                import json
                corners_dict = json.loads(clean_c)
                if isinstance(corners_dict, dict) and 'tl' in corners_dict and 'br' in corners_dict:
                    from app.image_preprocessor import warp_perspective_quad
                    warp_perspective_quad(dest_path, corners_dict, rotation=rotation, save_path=dest_path)
                    crop_applied = True
            except Exception as warp_err:
                print(f"[Upload] Notice during perspective unwarp: {warp_err}")
        elif rotation != 0:
            try:
                cv_img = cv2.imread(str(dest_path))
                if rotation == 90:
                    cv_img = cv2.rotate(cv_img, cv2.ROTATE_90_CLOCKWISE)
                elif rotation == 180:
                    cv_img = cv2.rotate(cv_img, cv2.ROTATE_180)
                elif rotation == 270:
                    cv_img = cv2.rotate(cv_img, cv2.ROTATE_90_COUNTERCLOCKWISE)
                cv2.imwrite(str(dest_path), cv_img)
            except Exception:
                pass

        # ── Step 2: Normalize EXIF orientation without applying duplicate enhancement ──
        if not crop_applied:
            try:
                with Image.open(dest_path) as im:
                    im = ImageOps.exif_transpose(im)
                    if im.mode not in ("RGB", "L"):
                        im = im.convert("RGB")
                    im.save(dest_path)
            except Exception:
                pass

        # Run chosen OCR engine
        ocr_lines, gemini_items = ocr_manager.extract_text(str(dest_path), engine=engine)
        
        # Match ingredients against database
        learning_result = None
        if gemini_items:
            ingredients = formula_parser.parse_gemini_items(gemini_items)
            learning_result = learning_engine.learn_from_google_scan(
                image_path=str(dest_path),
                local_ocr_lines=ocr_lines,
                gemini_verified_items=gemini_items,
                matched_database_items=ingredients
            )
        else:
            ingredients = formula_parser.parse_ocr_lines(ocr_lines)

        # If uploaded document had digital lines (PDF, Word, Excel, CSV, TXT), ensure nothing was missed
        if document_extracted_lines:
            doc_parsed = formula_parser.parse_ocr_lines([{"text": l, "confidence": 0.99, "box": []} for l in document_extracted_lines])
            if doc_parsed:
                if not gemini_items:
                    # When local OCR was used, digital document text is pristine vector text
                    final_ingredients = list(doc_parsed)
                    known_keys = {i.get("canonical_name", i.get("name", "")).lower().replace(" ", "") for i in final_ingredients}
                    for ocr_ing in ingredients:
                        c_k = ocr_ing.get("canonical_name", ocr_ing.get("name", "")).lower().replace(" ", "")
                        if c_k not in known_keys and not any(fuzz.ratio(c_k, kk) > 75 for kk in known_keys):
                            final_ingredients.append(ocr_ing)
                            known_keys.add(c_k)
                    ingredients = final_ingredients
                else:
                    existing_names = {i.get("name", "").lower().strip() for i in ingredients}
                    for p_item in doc_parsed:
                        p_name = p_item.get("name", "").lower().strip()
                        if p_name and p_name not in existing_names:
                            ingredients.append(p_item)
                            existing_names.add(p_name)
            
            # If OCR lines are sparse, synthesize OCR boxes from document text
            if (not ocr_lines or len(ocr_lines) < 2) and document_extracted_lines:
                ocr_lines = [
                    {"text": l, "confidence": 0.99, "box": [[25, idx * 36 + 135], [1175, idx * 36 + 135], [1175, idx * 36 + 167], [25, idx * 36 + 167]], "bbox": [25, idx * 36 + 135, 1175, idx * 36 + 167]}
                    for idx, l in enumerate(document_extracted_lines[:60])
                ]

        # Enrich ingredients with rate availability and Salt/RDA annotations
        software_rates = []
        for ing in ingredients:
            r_info = batch_master_engine.get_rate_info(ing.get("name", ""))
            ing["is_rate_available"] = r_info.get("available", False)
            ing["rate"] = r_info.get("rate", 0.0)
            ing["rate_matched_name"] = r_info.get("matched_name", "")
            
            # Annotate with Salt Element and RDA
            salt_rda_engine.annotate_ingredient(ing)

            software_rates.append({
                "name": ing.get("name"),
                "available": ing["is_rate_available"],
                "rate": ing["rate"],
                "matched_name": ing["rate_matched_name"]
            })

        # Scientific Evaluation
        evaluation = evaluator.evaluate_formula(ingredients)

        # Audit Log: Record upload and software rates in IST
        log_upload_event(
            username=username,
            image_url=f"/uploads/{unique_filename}",
            original_filename=orig_name,
            engine_used=engine,
            ingredients=ingredients,
            software_rates=software_rates
        )

        return JSONResponse({
            "success": True,
            "engine_used": engine,
            "image_url": f"/uploads/{unique_filename}",
            "ocr_lines": ocr_lines,
            "ingredients": ingredients,
            "matched_count": sum(1 for i in ingredients if i.get("is_matched")),
            "total_count": len(ingredients),
            "evaluation": evaluation,
            "learning_result": learning_result,
            "learning_status": learning_engine.get_learning_status()
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.post("/api/scan-sample")
async def scan_sample(
    request: Request,
    sample_name: str = Form(...),
    engine: str = Form("gemini")
):
    try:
        sample_path = SAMPLE_DIR / sample_name
        if not sample_path.exists():
            return JSONResponse({"success": False, "error": f"Sample image {sample_name} not found"}, status_code=404)

        ocr_lines, gemini_items = ocr_manager.extract_text(str(sample_path), engine=engine)

        learning_result = None
        if gemini_items:
            ingredients = formula_parser.parse_gemini_items(gemini_items)
            learning_result = learning_engine.learn_from_google_scan(
                image_path=str(sample_path),
                local_ocr_lines=ocr_lines,
                gemini_verified_items=gemini_items,
                matched_database_items=ingredients
            )
        else:
            ingredients = formula_parser.parse_ocr_lines(ocr_lines)

        for ing in ingredients:
            r_info = batch_master_engine.get_rate_info(ing.get("name", ""))
            ing["is_rate_available"] = r_info.get("available", False)
            ing["rate"] = r_info.get("rate", 0.0)
            ing["rate_matched_name"] = r_info.get("matched_name", "")
            salt_rda_engine.annotate_ingredient(ing)

        evaluation = evaluator.evaluate_formula(ingredients)

        return JSONResponse({
            "success": True,
            "engine_used": engine,
            "image_url": f"/samples/{sample_name}",
            "ocr_lines": ocr_lines,
            "ingredients": ingredients,
            "matched_count": sum(1 for i in ingredients if i.get("is_matched")),
            "total_count": len(ingredients),
            "evaluation": evaluation,
            "learning_result": learning_result,
            "learning_status": learning_engine.get_learning_status()
        })
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

# ══════════════════════════════════════════════════════════════════
# ELEMENTAL SALT & RDA CALCULATION ENGINE (PDF + INDEX (1).HTML)
# ══════════════════════════════════════════════════════════════════

@app.get("/api/salt-rda/catalog")
async def get_salt_rda_catalog():
    """
    Returns full database of minerals, salt forms, percentages, and RDA groups
    with convenient lists for frontend dropdowns
    """
    raw_cat = salt_rda_engine.get_full_catalog()
    minerals_dict = raw_cat.get("minerals", {})
    vitamins_dict = raw_cat.get("vitamins", {})
    misc_dict = raw_cat.get("miscellaneous", {})

    minerals_list = []
    salts_list = []
    for m_key, m_info in minerals_dict.items():
        m_name = m_info.get("name", m_key.title())
        minerals_list.append(m_name)
        for sf in m_info.get("saltForms", []):
            pct = float(sf.get("percentage") if sf.get("percentage") is not None else 1.0)
            salts_list.append({
                "element": m_name,
                "element_key": m_key,
                "salt_name": sf.get("name"),
                "elemental_percent": round(pct * 100, 2),
                "secondary": sf.get("secondary")
            })

    for misc_key, misc_info in misc_dict.items():
        misc_name = misc_info.get("name", misc_key.title())
        minerals_list.append(misc_name)
        for sf in misc_info.get("saltForms", []):
            pct = float(sf.get("percentage") if sf.get("percentage") is not None else 1.0)
            salts_list.append({
                "element": misc_name,
                "element_key": misc_key,
                "salt_name": sf.get("name"),
                "elemental_percent": round(pct * 100, 2),
                "secondary": sf.get("secondary")
            })

    vitamins_list = []
    vitamins_detail = {}
    for v_key, v_info in vitamins_dict.items():
        v_name = v_info.get("name", v_key.title())
        vitamins_list.append(v_name)
        vitamins_detail[v_name] = [f.get("name") for f in v_info.get("forms", [])]

    return JSONResponse({
        "success": True,
        "minerals": minerals_list,
        "salts": salts_list,
        "vitamins": vitamins_list,
        "vitamins_detail": vitamins_detail,
        "raw_catalog": raw_cat
    })

@app.post("/api/salt-rda/calculate")
async def calculate_salt_rda(request: Request):
    """
    Calculates elemental yield and demographic RDA breakdown table
    """
    try:
        body = await request.json()
        name = body.get("name") or body.get("salt_name") or ""
        dose = float(body.get("dose") or body.get("dose_mg") or 100)
        unit = body.get("unit", "mg")

        matched = salt_rda_engine.match_salt(name)
        if not matched:
            return JSONResponse({"success": False, "error": f"Salt or nutrient '{name}' not found in database"}, status_code=404)

        result = salt_rda_engine.calculate_elemental(matched, dose, unit)
        parent_key = matched.get("parent_key") or matched.get("parent_name")
        elem_options = salt_rda_engine.get_element_options(parent_key, result["elemental_amount_mg"], is_target_elemental=True, unit="mg").get("options", [])

        return JSONResponse({
            "success": True,
            "element_name": result["parent_nutrient"],
            "input_salt_name": result["salt_name"],
            "dose_mg": result["input_dose"],
            "unit": result["input_unit"],
            "elemental_mg": result["elemental_amount_mg"],
            "elemental_percent": result["salt_percentage"],
            "secondary_element": result["secondary"]["name"] if result.get("secondary") else None,
            "secondary_mg": result["secondary"]["amount_mg"] if result.get("secondary") else 0,
            "secondary_percent": result["secondary"]["percentage"] if result.get("secondary") else 0,
            "adult_rda_percent": result["standard_adult_rda_pct"],
            "demographic_rda": [
                {
                    "group": r["group"],
                    "rda_mg": r["rda_value"],
                    "unit": r["rda_unit"],
                    "pct": r["percentage"]
                }
                for r in result.get("rda_breakdown", [])
            ],
            "element_options": elem_options,
            "result": result
        })
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.post("/api/salt-rda/element-options")
async def get_salt_element_options(request: Request):
    """
    Returns all salt form options for an element with calculated required doses,
    elemental yields, secondary nutrients, and demographic RDA coverage.
    """
    try:
        body = await request.json()
        element = body.get("element") or body.get("name") or ""
        dose = float(body.get("dose") or body.get("target_dose") or body.get("dose_mg") or 100.0)
        is_target_elemental = bool(body.get("is_target_elemental", True))
        unit = body.get("unit", "mg")

        options_data = salt_rda_engine.get_element_options(
            text_or_key=element,
            dose=dose,
            is_target_elemental=is_target_elemental,
            unit=unit
        )
        if not options_data.get("success"):
            return JSONResponse(options_data, status_code=404)
        return JSONResponse(options_data)
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

# ══════════════════════════════════════════════════════════════════
# BATCH MASTER CARD (WITH PASSWORD PROTECTION & TABLET AUTO-SIZING)
# ══════════════════════════════════════════════════════════════════

@app.post("/api/batch-master/check-rates")
async def check_batch_master_rates(request: Request):
    try:
        body = await request.json()
        names = body.get("ingredients", [])
        results = batch_master_engine.batch_check_rates(names)
        return JSONResponse({"success": True, "results": results})
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.get("/api/batch-master/rates")
async def get_batch_master_rates(query: str = Query("", description="Search term for ingredients")):
    try:
        results = batch_master_engine.search_rates(query)
        return JSONResponse({"success": True, "results": results})
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.get("/api/batch-master/excipients")
async def get_batch_master_excipients(
    product_type: str = Query("Tablet"),
    tablet_size: int = Query(400)
):
    try:
        excipients = batch_master_engine.get_default_excipients(product_type, tablet_size)
        return JSONResponse({"success": True, "excipients": excipients})
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.post("/api/batch-master/calculate")
async def calculate_batch_master(request: Request):
    """
    Calculate theoretical & actual batch weights with strict auto-sizing and 1800mg tablet limit,
    and log final quotation rates to audit logs.
    """
    try:
        user = get_current_user(request)
        username = user["username"] if user else "guest"
        payload = await request.json()
        result = batch_master_engine.calculate_bmr(payload)

        # Audit log the exact rates given to user and track missing ingredient requests
        if result.get("success"):
            fin = result.get("financials", {})
            pkg_sum = result.get("packaging_summary", {})
            prim_type = pkg_sum.get("primary_type", "STRIP")
            if result.get("product_type") == "Liquid":
                rate_pack = fin.get("profit_per_bottle", 0.0)
                pack_name = f"Bottle ({result.get('size_or_capsule', '100ml')})"
            elif prim_type == "JAR":
                rate_pack = fin.get("profit_per_jar", 0.0)
                pack_name = f"Jar ({pkg_sum.get('num_jars', 0)} jars)"
            elif prim_type == "LOOSE":
                rate_pack = fin.get("profit_per_loose", 0.0)
                pack_name = f"Loose ({pkg_sum.get('num_loose', 0)} packs)"
            else:
                rate_pack = fin.get("profit_per_strip", 0.0)
                pack_name = f"Strip ({pkg_sum.get('num_strips', 0)} strips)"

            missing_items = fin.get("missing_rate_items", [])
            user_estimated_rates = {}
            for act in payload.get("active_ingredients", []):
                act_name = str(act.get("name", "")).strip()
                u_rate = float(act.get("user_estimated_rate") or (act.get("rate") if act.get("is_user_estimated") else 0.0) or 0.0)
                if act_name and u_rate > 0:
                    user_estimated_rates[act_name] = u_rate

            quote_id = log_quotation_event(
                username=username,
                product_type=result.get("product_type", "Tablet"),
                batch_qty=result.get("quantity", 100000),
                total_batch_kg=result.get("total_batch_kg", 0.0),
                status="QUOTATION_CALCULATED",
                pack_type=pack_name,
                rate_per_pack=rate_pack,
                rate_per_unit=fin.get("rate_per_unit", 0.0),
                total_batch_val=fin.get("total_batch_cost_with_profit", 0.0),
                missing_items=missing_items,
                raw_payload=payload,
                user_estimated_rates=user_estimated_rates
            )
            result["quotation_id"] = quote_id

        return JSONResponse(result)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.post("/api/batch-master/log-estimated-rate")
async def api_log_user_estimated_rate(request: Request):
    """
    Log user-filled estimated rate for an ingredient not in database.
    DOES NOT modify rate avg.xlsx, but maintains it in Admin audit panel.
    """
    try:
        user = get_current_user(request)
        username = user["username"] if user else "guest"
        body = await request.json()
        ing_name = str(body.get("ingredient_name", "")).strip()
        est_rate = float(body.get("estimated_rate", 0.0))
        prod_type = str(body.get("product_type", "Tablet"))
        notes = str(body.get("notes", "User-filled estimated rate"))

        if not ing_name or est_rate <= 0:
            return JSONResponse({"success": False, "error": "Invalid ingredient name or rate"}, status_code=400)

        req_id = log_user_estimated_rate(
            username=username,
            ingredient_name=ing_name,
            estimated_rate=est_rate,
            product_type=prod_type,
            notes=notes
        )
        return JSONResponse({
            "success": True,
            "request_id": req_id,
            "message": f"Estimated rate of ₹{est_rate:.2f}/kg recorded for Admin audit (Database unedited)."
        })
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.post("/api/batch-master/verify-password")
async def verify_bmr_master_password(request: Request):
    """
    Verify master password (Walpar@123) for unlocking or downloading the Batch Master Card
    """
    try:
        user = get_current_user(request)
        username = user["username"] if user else "guest"
        body = await request.json()
        password = str(body.get("password", "")).strip()

        if password == MASTER_BMR_PASSWORD:
            log_batch_download_event(
                username=username,
                product_type=body.get("product_type", "Tablet"),
                batch_qty=int(body.get("quantity", 100000)),
                total_batch_kg=float(body.get("total_batch_kg", 0.0)),
                status="UNLOCKED",
                password_used="***"
            )
            return JSONResponse({"success": True, "message": "Master Password Verified"})
        else:
            log_batch_download_event(
                username=username,
                product_type=body.get("product_type", "Tablet"),
                batch_qty=int(body.get("quantity", 100000)),
                total_batch_kg=float(body.get("total_batch_kg", 0.0)),
                status="PASSWORD_FAILED",
                password_used="***"
            )
            return JSONResponse({"success": False, "error": "Invalid Master Password. Access Denied."}, status_code=403)
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.post("/api/batch-master/export-excel")
async def export_batch_master_excel(request: Request):
    """
    Generate professional openpyxl Excel spreadsheet for Batch Master Card (Requires Walpar@123)
    """
    try:
        user = get_current_user(request)
        username = user["username"] if user else "guest"
        payload = await request.json()

        # Strict Master Password Check
        master_pass = str(payload.get("master_password", "")).strip()
        if master_pass != MASTER_BMR_PASSWORD:
            return JSONResponse({
                "success": False,
                "error": "Unauthorized: Master Password Walpar@123 is required to download official Batch Master Card"
            }, status_code=403)

        bmr_result = batch_master_engine.calculate_bmr(payload)
        if not bmr_result.get("success"):
            return JSONResponse(bmr_result, status_code=400)

        fin = bmr_result.get("financials", {})
        pkg_sum = bmr_result.get("packaging_summary", {})
        prim_type = pkg_sum.get("primary_type", "STRIP")
        if bmr_result.get("product_type") == "Liquid":
            rate_pack = fin.get("profit_per_bottle", 0.0)
            pack_name = f"Bottle ({bmr_result.get('size_or_capsule', '100ml')})"
        elif prim_type == "JAR":
            rate_pack = fin.get("profit_per_jar", 0.0)
            pack_name = f"Jar ({pkg_sum.get('num_jars', 0)} jars)"
        elif prim_type == "LOOSE":
            rate_pack = fin.get("profit_per_loose", 0.0)
            pack_name = f"Loose ({pkg_sum.get('num_loose', 0)} packs)"
        else:
            rate_pack = fin.get("profit_per_strip", 0.0)
            pack_name = f"Strip ({pkg_sum.get('num_strips', 0)} strips)"

        log_batch_download_event(
            username=username,
            product_type=str(payload.get("product_type", "Product")),
            batch_qty=int(payload.get("quantity", 100000)),
            total_batch_kg=float(bmr_result.get("total_batch_kg", 0.0)),
            status="DOWNLOADED_EXCEL",
            password_used="***",
            pack_type=pack_name,
            rate_per_pack=rate_pack,
            rate_per_unit=fin.get("rate_per_unit", 0.0),
            total_batch_val=fin.get("total_batch_cost_with_profit", 0.0)
        )

        excel_stream = batch_master_engine.generate_excel(bmr_result)

        product_name = str(payload.get("product_type", "Product")).lower()
        filename = f"Walpar_BMR_{product_name}_{payload.get('quantity', 100000)}.xlsx"

        return StreamingResponse(
            excel_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.post("/api/admin/rates/upload")
async def api_admin_upload_rates(
    request: Request,
    file: UploadFile = File(...)
):
    """
    Upload rate sheet (.xlsx, .csv). Update existing rates and intelligently match similar ingredients.
    """
    user = get_current_user(request)
    if not user or user.get("role") != "admin":
        return JSONResponse({"success": False, "error": "Unauthorized: Administrator access required."}, status_code=403)
    try:
        content = await file.read()
        res = batch_master_engine.update_rates_from_file(content, file.filename or "rates.xlsx")
        return JSONResponse(res)
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.get("/api/admin/rates/list")
async def api_admin_rates_list(request: Request):
    """
    Returns list of all active ingredient rates for the Admin view.
    """
    user = get_current_user(request)
    if not user or user.get("role") != "admin":
        return JSONResponse({"success": False, "error": "Unauthorized"}, status_code=403)
    rates = batch_master_engine.get_all_ingredient_rates()
    return JSONResponse({"success": True, "count": len(rates), "rates": rates})

@app.post("/api/admin/rates/update-single")
async def api_admin_update_single_rate(request: Request):
    """
    Directly update or add an ingredient rate from admin panel,
    and automatically recalculate any pending user quotations!
    """
    user = get_current_user(request)
    if not user or user.get("role") != "admin":
        return JSONResponse({"success": False, "error": "Unauthorized"}, status_code=403)
    try:
        body = await request.json()
        ingredient_name = body.get("name", "").strip()
        new_rate = float(body.get("rate", 0.0))
        if not ingredient_name:
            return JSONResponse({"success": False, "error": "Ingredient name required"}, status_code=400)
        if new_rate < 0:
            return JSONResponse({"success": False, "error": "Rate cannot be negative"}, status_code=400)

        # 1. Update in master database & excel
        rate_res = batch_master_engine.update_single_rate(ingredient_name, new_rate)
        if not rate_res.get("success"):
            return JSONResponse(rate_res, status_code=400)

        # 2. Check and resolve any pending requests
        affected_quote_ids = resolve_missing_rate_requests(ingredient_name, new_rate, user["username"])
        
        # 3. Recalculate affected quotations
        recalculated_quotes = []
        for qid in affected_quote_ids:
            q_info = get_quotation_by_id(qid)
            if q_info and q_info.get("raw_payload"):
                payload = q_info["raw_payload"]
                recalc = batch_master_engine.calculate_bmr(payload)
                if recalc.get("success"):
                    r_fin = recalc.get("financials", {})
                    r_pkg = recalc.get("packaging_summary", {})
                    prim = r_pkg.get("primary_type", "STRIP")
                    if recalc.get("product_type") == "Liquid":
                        r_pack = r_fin.get("profit_per_bottle", 0.0)
                    elif prim == "JAR":
                        r_pack = r_fin.get("profit_per_jar", 0.0)
                    elif prim == "LOOSE":
                        r_pack = r_fin.get("profit_per_loose", 0.0)
                    else:
                        r_pack = r_fin.get("profit_per_strip", 0.0)

                    remaining_missing = r_fin.get("missing_rate_items", [])
                    note = f"Admin updated '{ingredient_name}' to ₹{new_rate:.2f}/Kg on {get_current_ist_str()}"
                    update_quotation_revised_rates(
                        quotation_id=qid,
                        new_rate_pack=r_pack,
                        new_rate_unit=r_fin.get("rate_per_unit", 0.0),
                        new_total_val=r_fin.get("total_batch_cost_with_profit", 0.0),
                        remaining_missing=remaining_missing,
                        note=note
                    )
                    recalculated_quotes.append({
                        "quotation_id": qid,
                        "username": q_info.get("username"),
                        "product_type": q_info.get("product_type"),
                        "revised_rate_per_pack": r_pack,
                        "revised_total_val": r_fin.get("total_batch_cost_with_profit", 0.0)
                    })

        return JSONResponse({
            "success": True,
            "rate_update": rate_res,
            "resolved_requests_count": len(affected_quote_ids),
            "recalculated_quotes": recalculated_quotes,
            "message": f"Successfully updated '{rate_res.get('ingredient')}' to ₹{new_rate:,.2f}/Kg. {len(recalculated_quotes)} user quotation(s) updated with new rates."
        })
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.post("/api/admin/rates/resolve-missing")
async def api_admin_resolve_missing_rate(request: Request):
    """
    Explicit endpoint for admin to set the missing ingredient price and push revised quotation to user
    """
    user = get_current_user(request)
    if not user or user.get("role") != "admin":
        return JSONResponse({"success": False, "error": "Unauthorized"}, status_code=403)
    try:
        body = await request.json()
        ingredient_name = body.get("ingredient_name", "").strip()
        new_rate = float(body.get("rate", 0.0))
        if not ingredient_name or new_rate <= 0:
            return JSONResponse({"success": False, "error": "Valid ingredient name and rate (>0) required"}, status_code=400)

        # Update rate in master database
        rate_res = batch_master_engine.update_single_rate(ingredient_name, new_rate)
        if not rate_res.get("success"):
            return JSONResponse(rate_res, status_code=400)

        # Resolve pending requests
        affected_quote_ids = resolve_missing_rate_requests(ingredient_name, new_rate, user["username"])

        # Recalculate affected quotations
        recalculated_quotes = []
        for qid in affected_quote_ids:
            q_info = get_quotation_by_id(qid)
            if q_info and q_info.get("raw_payload"):
                payload = q_info["raw_payload"]
                recalc = batch_master_engine.calculate_bmr(payload)
                if recalc.get("success"):
                    r_fin = recalc.get("financials", {})
                    r_pkg = recalc.get("packaging_summary", {})
                    prim = r_pkg.get("primary_type", "STRIP")
                    if recalc.get("product_type") == "Liquid":
                        r_pack = r_fin.get("profit_per_bottle", 0.0)
                    elif prim == "JAR":
                        r_pack = r_fin.get("profit_per_jar", 0.0)
                    elif prim == "LOOSE":
                        r_pack = r_fin.get("profit_per_loose", 0.0)
                    else:
                        r_pack = r_fin.get("profit_per_strip", 0.0)

                    remaining_missing = r_fin.get("missing_rate_items", [])
                    note = f"Missing ingredient '{ingredient_name}' priced at ₹{new_rate:.2f}/Kg by Admin ({user['username']}) on {get_current_ist_str()}"
                    update_quotation_revised_rates(
                        quotation_id=qid,
                        new_rate_pack=r_pack,
                        new_rate_unit=r_fin.get("rate_per_unit", 0.0),
                        new_total_val=r_fin.get("total_batch_cost_with_profit", 0.0),
                        remaining_missing=remaining_missing,
                        note=note
                    )
                    recalculated_quotes.append({
                        "quotation_id": qid,
                        "username": q_info.get("username"),
                        "product_type": q_info.get("product_type"),
                        "revised_rate_per_pack": r_pack,
                        "revised_total_val": r_fin.get("total_batch_cost_with_profit", 0.0)
                    })

        return JSONResponse({
            "success": True,
            "message": f"Added rate ₹{new_rate:,.2f}/Kg for '{ingredient_name}'. {len(recalculated_quotes)} user quotation(s) recalculated and updated.",
            "recalculated_quotes": recalculated_quotes
        })
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)

@app.get("/api/admin/missing-rates")
async def api_admin_missing_rates(request: Request):
    """
    Returns pending and resolved missing rate requests for the admin notifications panel
    """
    user = get_current_user(request)
    if not user or user.get("role") != "admin":
        return JSONResponse({"success": False, "error": "Unauthorized"}, status_code=403)
    pending = get_pending_missing_rates()
    all_requests = get_all_missing_rate_requests(limit=100)
    return JSONResponse({
        "success": True,
        "pending_count": len(pending),
        "pending": pending,
        "all_requests": all_requests
    })

@app.get("/api/admin/user-quotation-history")
async def api_admin_user_history(
    request: Request,
    username: Optional[str] = Query(None)
):
    """
    Returns per-user quotation history and aggregate portfolio statistics
    """
    user = get_current_user(request)
    if not user or user.get("role") != "admin":
        return JSONResponse({"success": False, "error": "Unauthorized"}, status_code=403)
    quotes = get_batch_logs(limit=200, user_filter=username)
    per_user_stats = get_per_user_quotation_stats()
    return JSONResponse({
        "success": True,
        "filter_user": username,
        "count": len(quotes),
        "quotations": quotes,
        "per_user_stats": per_user_stats
    })

@app.get("/api/user/quotations")
async def api_get_my_quotations(request: Request):
    """
    Returns current user's quotation history with updated rates from admin
    """
    user = get_current_user(request)
    if not user:
        return JSONResponse({"success": False, "error": "Authentication required"}, status_code=401)
    quotes = get_batch_logs(limit=100, user_filter=user["username"])
    # Check if any quotation has updated rate from admin
    has_updated_quotes = any(q.get("status") == "RATE_UPDATED_BY_ADMIN" or q.get("rate_status") == "RESOLVED_BY_ADMIN" for q in quotes)
    return JSONResponse({
        "success": True,
        "username": user["username"],
        "count": len(quotes),
        "has_updated_quotes": has_updated_quotes,
        "quotations": quotes
    })

@app.get("/api/admin/master-config")
async def api_admin_get_master_config(request: Request):
    user = get_current_user(request)
    if not user or user.get("role") != "admin":
        return JSONResponse({"success": False, "error": "Unauthorized"}, status_code=403)
    return JSONResponse({"success": True, "config": master_config_mgr.get_config()})

@app.post("/api/admin/master-config/update")
async def api_admin_update_master_config(request: Request):
    user = get_current_user(request)
    if not user or user.get("role") != "admin":
        return JSONResponse({"success": False, "error": "Unauthorized"}, status_code=403)
    try:
        body = await request.json()
        updated = master_config_mgr.update_config(body)
        batch_master_engine._load_liquid_ingredients()
        return JSONResponse({"success": True, "config": updated, "message": "Master packaging & production configuration successfully updated!"})
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)


